from __future__ import annotations

import sys
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.errors import BadRequestError, ConvertError
from app.skills.tuoshu import convert_service, converter


def _make_complex_docx(tmp_path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (20, 20), color="white").save(image_path)

    source = tmp_path / "complex.docx"
    document = Document()
    document.add_paragraph("正文提单号：TEST001")
    table = document.add_table(rows=1, cols=3)
    merged_cell = table.cell(0, 0).merge(table.cell(0, 2))
    merged_cell.text = "Word 合并标题"
    document.sections[0].header.paragraphs[0].text = "页眉客户编号：HEADER001"
    document.sections[0].footer.paragraphs[0].text = "页脚联系人：张三"
    document.add_picture(str(image_path), width=Inches(0.2))
    document.save(source)

    with zipfile.ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    document_xml = members["word/document.xml"]
    text_box_xml = (
        b'<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml">'
        b"<v:textbox><w:txbxContent><w:p><w:r>"
        b"<w:t>\xe6\x96\x87\xe6\x9c\xac\xe6\xa1\x86\xe4\xb8\x9a\xe5\x8a\xa1\xe5\x8f\xb7\xef\xbc\x9aTB001</w:t>"
        b"</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    members["word/document.xml"] = document_xml.replace(
        b"<w:sectPr", text_box_xml + b"<w:sectPr", 1
    )
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return source


def test_docx_story_text_textboxes_and_images_are_discovered(tmp_path):
    source = _make_complex_docx(tmp_path)

    output, report = converter.convert_docx(str(source))

    assert "正文提单号：TEST001" in output
    assert output.count("Word 合并标题") == 1
    assert "页眉客户编号：HEADER001" in output
    assert "页脚联系人：张三" in output
    assert "文本框业务号：TB001" in output
    assert "## Embedded images" in output
    assert report["coverage"]["story_parts"] >= 3
    assert report["coverage"]["story_parts_converted"] == report["coverage"]["story_parts"]
    assert report["coverage"]["text_boxes"] == 1
    assert report["coverage"]["embedded_images"] == 1
    assert len(report["embedded_images"]) == 1


def test_docx_embedded_image_is_routed_as_visual_evidence(tmp_path, monkeypatch):
    from app.config import settings

    source = _make_complex_docx(tmp_path)
    monkeypatch.setattr(settings, "mineru_enabled", False)

    converted = convert_service.convert_to_markdown(source.read_bytes(), source.name)
    parse_result = converted.parse_result

    assert parse_result.parser == "mixed"
    assert len(parse_result.vision_inputs) == 1
    assert parse_result.vision_inputs[0][1] == "image/png"
    assert parse_result.coverage["embedded_images"] == 1
    assert parse_result.coverage["embedded_images_processed"] == 1
    issue_codes = {issue["code"] for issue in parse_result.review_issues()}
    assert "vision_only_unverified" in issue_codes
    assert "embedded_image_requires_review" in issue_codes


def test_xlsx_merged_cells_are_not_duplicated_and_dates_use_display_precision(tmp_path):
    import openpyxl

    source = tmp_path / "merged.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "做箱通知书"
    sheet.merge_cells("A1:C1")
    sheet["A2"] = "做箱时间"
    sheet["B2"] = datetime(2022, 1, 20)
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet.merge_cells("B2:C2")
    workbook.save(source)

    converted = convert_service.convert_to_markdown(source.read_bytes(), source.name)

    assert str(converted).count("做箱通知书") == 1
    assert "| 1 | 做箱通知书 |  |" in converted
    assert "| 2 | 做箱时间 | 2022-01-20 |" in converted
    assert "2022-01-20 00:00:00" not in converted
    assert "- `A1:C1`" in converted
    assert "- `B2:C2`" in converted
    assert converted.parse_result.coverage == {
        "sheet_count": 1,
        "nonempty_cells": 3,
        "merged_ranges": 2,
        "formula_cells": 0,
        "formula_values_available": 0,
        "complete": True,
        "omissions": [],
    }


def test_xlsx_formula_and_missing_cached_value_are_explicit(tmp_path):
    import openpyxl

    source = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = 2
    sheet["A2"] = 3
    sheet["A3"] = "=SUM(A1:A2)"
    workbook.save(source)

    converted = convert_service.convert_to_markdown(source.read_bytes(), source.name)

    assert "| A3 | =SUM(A1:A2) | _unavailable_ |" in converted
    assert converted.parse_result.coverage["formula_cells"] == 1
    assert converted.parse_result.coverage["formula_values_available"] == 0
    assert converted.parse_result.coverage["complete"] is False
    issue = converted.parse_result.review_issues()[0]
    assert issue["code"] == "formula_value_unavailable"
    assert issue["source_values"] == ["Sheet!A3"]
    assert issue["blocking"] is True


def test_word_2003_xml_is_converted_without_external_tool(monkeypatch, tmp_path):
    source = tmp_path / "wordml.doc"
    source.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<w:wordDocument xmlns:w="http://schemas.microsoft.com/office/word/2003/wordml">
  <w:body>
    <w:p><w:r><w:t>配舱通知</w:t></w:r></w:p>
    <w:tbl><w:tr>
      <w:tc><w:p><w:r><w:t>提单号</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:t>SITG001</w:t></w:r></w:p></w:tc>
    </w:tr></w:tbl>
  </w:body>
</w:wordDocument>""",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        converter,
        "_find_soffice",
        lambda: (_ for _ in ()).throw(AssertionError("office lookup should not run")),
    )

    output, _ = converter.convert_doc(str(source))

    assert "_format: word_xml_" in output
    assert "_p1_ 配舱通知" in output
    assert "| 1 | 提单号 | SITG001 |" in output


def test_flat_opc_word_xml_extracts_document_body(monkeypatch, tmp_path):
    source = tmp_path / "flat-opc.doc"
    source.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<pkg:package xmlns:pkg="http://schemas.microsoft.com/office/2006/xmlPackage"
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <pkg:part pkg:name="/word/document.xml"><pkg:xmlData>
    <w:document><w:body><w:p><w:r><w:t>放箱号 0156</w:t></w:r></w:p></w:body></w:document>
  </pkg:xmlData></pkg:part>
</pkg:package>""",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        converter,
        "_find_soffice",
        lambda: (_ for _ in ()).throw(AssertionError("office lookup should not run")),
    )

    output, _ = converter.convert_doc(str(source))

    assert "_format: word_xml_" in output
    assert "_p1_ 放箱号 0156" in output
    assert "pkg:package" not in output


def test_doc_uses_textutil_when_libreoffice_is_unavailable(monkeypatch, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-word")
    captured: dict[str, object] = {}

    monkeypatch.setattr(converter, "_find_soffice", lambda: None)
    monkeypatch.setattr(converter, "_find_textutil", lambda: "/usr/bin/textutil")

    def fake_run(command, **kwargs):
        captured["command"] = command
        captured["kwargs"] = kwargs
        Path(command[command.index("-output") + 1]).write_bytes(b"converted-docx")

    monkeypatch.setattr(converter.subprocess, "run", fake_run)
    monkeypatch.setattr(
        converter,
        "convert_docx",
        lambda path: (captured.update(converted_path=path), ("", {}))[1],
    )

    converter.convert_doc(str(source))

    command = captured["command"]
    assert command[:3] == ["/usr/bin/textutil", "-convert", "docx"]
    assert command[-2:] == ["--", str(source)]
    assert Path(captured["converted_path"]).name == "legacy.docx"
    assert captured["kwargs"]["timeout"] == 120


def test_doc_textutil_output_uses_ascii_filename(monkeypatch, tmp_path):
    source = tmp_path / "2×40HQ余姚-上海排柜托书.doc"
    source.write_bytes(b"legacy-word")
    captured: dict[str, object] = {}

    monkeypatch.setattr(converter, "_find_soffice", lambda: None)
    monkeypatch.setattr(converter, "_find_textutil", lambda: "/usr/bin/textutil")

    def fake_run(command, **kwargs):
        captured["command"] = command
        out = Path(command[command.index("-output") + 1])
        captured["out_name"] = out.name
        out.write_bytes(b"converted-docx")

    monkeypatch.setattr(converter.subprocess, "run", fake_run)
    monkeypatch.setattr(converter, "convert_docx", lambda _path: ("", {}))

    converter.convert_doc(str(source))

    out_name = str(captured["out_name"])
    assert all(char.isascii() for char in out_name)
    assert out_name.endswith(".docx")


def test_make_temp_dir_prefers_tmpdir_env(monkeypatch, tmp_path):
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    monkeypatch.setattr(converter.tempfile, "gettempdir", lambda: "/nonexistent-system-tmp")

    with converter._make_temp_dir(prefix="probe-") as tmpd:
        assert str(tmpd).startswith(str(tmp_path))
        assert Path(tmpd).exists()
        assert Path(tmpd).name.startswith("probe-")


def test_make_temp_dir_falls_back_when_tmpdir_unwritable(monkeypatch, tmp_path):
    blocked = tmp_path / "blocked"
    blocked.mkdir()
    blocked.chmod(0o500)
    monkeypatch.setenv("TMPDIR", str(blocked))
    try:
        with converter._make_temp_dir(prefix="probe-") as tmpd:
            assert Path(tmpd).exists()
            assert not str(tmpd).startswith(str(blocked))
    finally:
        blocked.chmod(0o700)


def test_doc_prefers_libreoffice_when_both_converters_exist(monkeypatch, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-word")
    captured: dict[str, object] = {}

    monkeypatch.setattr(converter, "_find_soffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(
        converter,
        "_find_textutil",
        lambda: (_ for _ in ()).throw(AssertionError("textutil lookup should not run")),
    )

    def fake_run(command, **_kwargs):
        captured["command"] = command
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "legacy.docx").write_bytes(b"converted-docx")

    monkeypatch.setattr(converter.subprocess, "run", fake_run)
    monkeypatch.setattr(converter, "convert_docx", lambda _path: ("", {}))

    converter.convert_doc(str(source))

    assert captured["command"][:2] == ["/usr/bin/soffice", "--headless"]


def test_doc_without_local_converter_emits_actionable_hint(monkeypatch, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-word")
    monkeypatch.setattr(converter, "_find_soffice", lambda: None)
    monkeypatch.setattr(converter, "_find_textutil", lambda: None)

    output, _ = converter.convert_doc(str(source))

    assert output.startswith("SCAN_OR_IMAGE_HINT:")
    assert "LibreOffice" in output
    assert "textutil" in output


def test_concurrent_local_conversions_do_not_mix_outputs(monkeypatch):
    def fake_converter(path):
        name = Path(path).name
        time.sleep(0.03)
        return f"{name}:start\n{name}:end\n", {}

    monkeypatch.setitem(convert_service._DISPATCH, ".docx", fake_converter)
    monkeypatch.setattr(convert_service, "validate_document_content", lambda *_args: "docx")

    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(convert_service.convert_to_markdown, b"a", "first.docx")
        second = pool.submit(convert_service.convert_to_markdown, b"b", "second.docx")
        outputs = {str(first.result()), str(second.result())}

    assert outputs == {
        "first.docx:start\nfirst.docx:end\n",
        "second.docx:start\nsecond.docx:end\n",
    }


def test_document_extension_mismatch_is_rejected():
    with pytest.raises(BadRequestError) as exc_info:
        convert_service.validate_document_content(b"%PDF-1.7\n", "renamed.docx")

    assert exc_info.value.code == "file_format_mismatch"
    assert exc_info.value.details["detected_format"] == "pdf"


def test_empty_converter_output_is_rejected(monkeypatch):
    monkeypatch.setattr(convert_service, "validate_document_content", lambda *_args: "docx")
    monkeypatch.setitem(convert_service._DISPATCH, ".docx", lambda _path: ("", {}))

    with pytest.raises(ConvertError) as exc_info:
        convert_service.convert_to_markdown(b"document", "empty.docx")

    assert exc_info.value.code == "empty_converted_content"


def test_pdf_vision_conversion_rejects_page_truncation(monkeypatch):
    class FakePdf:
        def __len__(self):
            return 4

        def close(self):
            return None

    fake_pdfium = SimpleNamespace(PdfDocument=lambda _bytes: FakePdf())
    monkeypatch.setitem(sys.modules, "pypdfium2", fake_pdfium)

    with pytest.raises(ConvertError) as exc_info:
        convert_service.render_pdf_pages(b"%PDF-1.7\n", max_pages=3, scale=2.0)

    assert exc_info.value.code == "pdf_page_limit_exceeded"
    assert exc_info.value.details == {"page_count": 4, "max_pages": 3}
