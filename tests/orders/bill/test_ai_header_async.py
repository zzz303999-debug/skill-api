"""两段式 L3 表头映射测试（生产编排链路，mock achat_json，零网络）。

收尾改造（docs/异步化改造收尾计划.md 改造项一）：生产解析编排
_parse_stage_async —— L1/L2 命中路径零 LLM 调用；L3 未命中 →
AiHeaderNeeded 信号 → service.achat_json（真异步）→ 闸门/解析二次进段。
与 test_template.py 的 L3 语义用例同口径（均已走 _parse_stage_async 两段式）
语义对齐；本文件 patch 点为 service 命名空间 achat_json（编排层 import）。
"""

from __future__ import annotations

import pytest
from openpyxl.utils import get_column_letter

import app.orders.bill.service as service_module
from app.core.errors import BadRequestError, LLMError
from app.orders.bill.parsing.opener import AiHeaderNeeded, open_and_identify
from app.orders.bill.parsing.parser import ParseOutput
from helpers import build_bill_bytes

pytestmark = pytest.mark.asyncio

# 异构模板（列名全换，模板库不命中）：与 test_template.HETERO_* 同构
HETERO_HEADERS = {
    "A": "序号",
    "B": "海运提单号",
    "C": "委托单位",
    "D": "箱型规格",
    "E": "海运费",
}
HETERO_ROWS = [
    {"A": 1, "B": "OOLU12345678", "C": "客户甲", "D": "40HQ", "E": 100},
    {"A": 2, "B": "OOLU12345679", "C": "客户乙", "D": "20GP", "E": 200},
]

# mock AI 正确映射（build_bill_bytes 表头在第 2 行，5 列）
AI_MAPPING_OK = {
    "header_row": 2,
    "two_row": False,
    "fee_boundary": "column_range",
    "mapping": [
        {"col": 1, "target": "ignore"},
        {"col": 2, "target": "bl_no"},
        {"col": 3, "target": "customer_name"},
        {"col": 4, "target": "box_type_qty"},
        {"col": 5, "target": "fee:运费"},
    ],
    "confidence": 0.95,
}

# 缺必映射字段 box_type_qty（闸门 a 拒绝场景，与 test_route 同款）
AI_MAPPING_MISSING_BOX = {
    "header_row": 2,
    "two_row": False,
    "fee_boundary": "column_range",
    "mapping": [
        {"col": 1, "target": "ignore"},
        {"col": 2, "target": "bl_no"},
        {"col": 3, "target": "customer_name"},
        {"col": 5, "target": "fee:运费"},
    ],
    "confidence": 0.9,
}


def _patch_async_ai(
    monkeypatch, payload: dict | None = None, exc: Exception | None = None
) -> dict:
    """mock service.achat_json（生产两段式 LLM 入口），计数调用次数。"""
    calls = {"n": 0}

    async def fake_achat_json(messages, **_kwargs):
        calls["n"] += 1
        if exc is not None:
            raise exc
        return dict(payload), {"model": "fake", "usage": None}

    monkeypatch.setattr(service_module, "achat_json", fake_achat_json)
    return calls


def _hetero_bytes() -> bytes:
    return build_bill_bytes(HETERO_HEADERS, HETERO_ROWS)


class TestOpenAndIdentify:
    """阶段 1（open_and_identify）：命中/未命中分流与资源责任移交。"""

    def test_unmatched_returns_needed_with_view(self, tmp_path):
        """异构模板未命中 → AiHeaderNeeded（view 跨段可用 + zone_lines 就绪）。"""
        path = tmp_path / "hetero.xlsx"
        path.write_bytes(_hetero_bytes())
        needed = open_and_identify(path)
        assert isinstance(needed, AiHeaderNeeded)
        assert needed.engine == "openpyxl"
        assert needed.filename == "hetero.xlsx"
        # 表头区文本（第一行行首 1-based 行号）与 view 跨段可读（第二段依赖）
        assert needed.zone_lines[0].startswith("1 |")
        assert needed.view.cell(2, 2) == "海运提单号"
        needed.close()
        needed.close()  # 幂等

    def test_builtin_hit_returns_output_without_llm(self, tmp_path, monkeypatch):
        """内置模板指纹命中 → 单段完成返回 ParseOutput（零 LLM，常见路径）。"""
        from app.orders.bill.parsing.legacy_template import BUILTIN_HEADERS

        calls = _patch_async_ai(monkeypatch)  # 未抛异常即未调用；计数校验见下
        # 32 列全量构造（超出 A-Z 用 AA.. 列字母，指纹按全列拼接计算）
        headers = {
            get_column_letter(i): h for i, h in enumerate(BUILTIN_HEADERS, start=1)
        }
        path = tmp_path / "builtin.xlsx"
        path.write_bytes(build_bill_bytes(headers, [{"B": 1}]))
        out = open_and_identify(path)
        assert isinstance(out, ParseOutput)
        assert not isinstance(out, AiHeaderNeeded)
        assert out.template is not None and out.template["source"] == "builtin"
        assert calls["n"] == 0


class TestTwoStageOrchestration:
    """两段式编排（_parse_stage_async）：L3 成功 / 回退 / 闸门拒绝。"""

    async def test_l3_mapped_and_aggregated(self, tmp_path, monkeypatch):
        """L3 未命中 → achat_json 真异步映射 → 解析成功 + 候选模板（对齐同步语义）。"""
        calls = _patch_async_ai(monkeypatch, AI_MAPPING_OK)
        output = await service_module._parse_stage_async("hetero.xlsx", _hetero_bytes())
        assert calls["n"] == 1
        assert isinstance(output, ParseOutput)
        assert output.template is not None
        assert output.template["source"] == "template"
        # 断言口径与同步链路 test_template.test_hetero_mapped_and_aggregated
        # 一致：L3 模板路径产 canonical_rows（BillRow 管线不产出）
        assert output.canonical_rows is not None
        assert len(output.canonical_rows) == 2
        assert output.canonical_rows[0]["bl_no"] == "OOLU12345678"
        assert output.new_template is not None  # 候选模板配置（预览固化用）
        assert output.new_template["columns"]["bl_no"] == ["海运提单号"]

    async def test_llm_error_falls_back_exact(self, tmp_path, monkeypatch):
        """LLM 不可用 → 回退精确匹配 → 异构表头无表头行 → 400 bad_request。

        与同步链路（test_route.test_llm_unavailable_falls_back_exact_400）
        同语义：LLM 故障降级而非 500。
        """
        _patch_async_ai(monkeypatch, exc=LLMError("LLM unavailable in tests"))
        with pytest.raises(BadRequestError) as caught:
            await service_module._parse_stage_async("hetero.xlsx", _hetero_bytes())
        assert caught.value.http_status == 400
        assert caught.value.code == "bad_request"
        assert "required_headers" in caught.value.details
        assert "missing_headers" in caught.value.details

    async def test_gate_rejected_header_mapping_rejected(self, tmp_path, monkeypatch):
        """AI 映射缺必映射字段 → 闸门拒绝 → 400 header_mapping_rejected。"""
        calls = _patch_async_ai(monkeypatch, AI_MAPPING_MISSING_BOX)
        with pytest.raises(BadRequestError) as caught:
            await service_module._parse_stage_async("hetero.xlsx", _hetero_bytes())
        assert calls["n"] == 1
        assert caught.value.http_status == 400
        assert caught.value.code == "header_mapping_rejected"
        details = caught.value.details
        assert any("box_type_qty" in reason for reason in details["failure_reasons"])
        assert details["ai_mapping"]["mapping"][1]["target"] == "bl_no"
        assert details["header_zone"][0].startswith("1 |")

    async def test_builtin_hit_zero_llm_through_orchestration(
        self, tmp_path, monkeypatch
    ):
        """常见路径（L1 命中）经完整编排零 LLM 调用、结果即第一段输出。"""
        from app.orders.bill.parsing.legacy_template import BUILTIN_HEADERS

        calls = _patch_async_ai(monkeypatch)
        headers = {
            get_column_letter(i): h for i, h in enumerate(BUILTIN_HEADERS, start=1)
        }
        output = await service_module._parse_stage_async(
            "builtin.xlsx", build_bill_bytes(headers, [{"B": 1}])
        )
        assert calls["n"] == 0
        assert output.template is not None
        assert output.template["source"] == "builtin"
