"""竞品账单测试共享工具与 golden 常量（pytest rootdir 模式：目录在 sys.path，直接 import）。"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import xlwt
from openpyxl import Workbook

from app.orders.bill import BillPeriod, group_orders

GOLDEN_DIR = Path(__file__).resolve().parent.parent.parent / "golden" / "bill"
REAL_XLS = GOLDEN_DIR / "2015-01到2015-12上海通寰应收对账单.xls"
REBUILT_XLSX = GOLDEN_DIR / "2015-01到2015-12上海通寰应收对账单-rebuilt.xlsx"
PERIOD_2015 = BillPeriod(start="2015-01-01", end="2015-12-31")

# 真实账单基线（2026-08 探查确认）：1094 原始行 → 尾部 4 行过滤 → 1090 数据行；
# 一行一票（2026-08-31 业务拍板）：单数 = 数据行数
REAL_RAW_ROWS = 1094
REAL_TOTAL_ROWS = 1090
REAL_ORDER_COUNT = REAL_TOTAL_ROWS


class FakeResponse:
    """下游 HTTP 响应替身：payload 序列化为 text；json() 返回 payload。"""

    def __init__(self, payload, status_code: int = 200):
        self._payload = payload
        self.status_code = status_code
        self.text = (
            json.dumps(payload, ensure_ascii=False)
            if isinstance(payload, (dict, list))
            else str(payload)
        )
        self.headers = {"content-type": "application/json"}

    def json(self, **kwargs):
        # 真实 httpx 行为：参数透传 json.loads（如 parse_constant）；
        # 无参数时直接返回预置 payload（既有测试零变化）
        if kwargs:
            return json.loads(self.text, **kwargs)
        return self._payload


def cells_equal(a, b) -> bool:
    """字段值等价：数值归一（'1' vs '1.0'）比较，文本严格相等。"""
    if a is None and b is None:
        return True
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return a == b


def build_bill_bytes(headers: dict[str, str], rows: list[dict[str, object]]) -> bytes:
    """内存构造账单：headers 列名 → 列字母（第 2 行为表头）；rows 每行 {列字母: 值}。"""
    wb = Workbook()
    ws = wb.active
    for col_letter, header in headers.items():
        ws[f"{col_letter}2"] = header
    for i, row in enumerate(rows, start=3):
        for col_letter, value in row.items():
            ws[f"{col_letter}{i}"] = value
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_bill_xls_bytes(headers: dict[str, str], rows: list[dict[str, object]]) -> bytes:
    """内存构造 .xls 账单（xlwt）：布局与 build_bill_bytes 一致（第 2 行为表头）。

    供小行数越界等需要真实 .xls 引擎（xlrd）的用例使用。
    """
    wb = xlwt.Workbook()
    ws = wb.add_sheet("bill")
    for col_letter, header in headers.items():
        ws.write(1, _col_index(col_letter), header)
    for i, row in enumerate(rows, start=2):
        for col_letter, value in row.items():
            ws.write(i, _col_index(col_letter), value)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _col_index(letter: str) -> int:
    """列字母（A/B/C…）→ 0-based 列号。"""
    idx = 0
    for ch in letter.upper():
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx - 1


def go(rows, period):
    """group_orders 便捷包装：只取 orders（reconciliation 由对账测试单独验证）。"""
    return group_orders(rows, period).orders
