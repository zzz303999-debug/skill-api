"""路由层测试：preview/create 语义、sk 校验、重复上传照常 200、错误码。

2026-09-01 起响应统一外壳 {code, msg, data}（对齐账单录入口径）：业务数据
在 data 内，错误场景 msg 为可直接展示的中文错误信息。
"""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.orders.manifest.service as service_module
from app.main import app

CREATE_HEADERS = {"sk": "tk-test"}


def _post(
    client: TestClient,
    content: bytes,
    filename: str = "a.xlsx",
    data: dict | None = None,
    headers: dict | None = None,
):
    return client.post(
        "/orders/manifest/import",
        files={"file": (filename, content, "application/octet-stream")},
        data=data or {},
        headers=headers or {},
    )


class TestRoutePreview:
    """preview：200 + 统一外壳 + 解析结果；零下游调用。"""

    def test_preview_returns_parsed(self, auth_bytes):
        client = TestClient(app)
        r = _post(client, auth_bytes)
        assert r.status_code == 200
        body = r.json()
        # 统一外壳：code/msg/data（2026-09-01 起对齐账单录入口径）
        assert body["code"] == "200"
        assert body["msg"] == "请求成功"
        data = body["data"]
        assert data["create_order"] is False
        assert data["summary"] is None
        assert data["upstream"] is None
        order = data["orders"][0]
        assert order["bl_no"] == "SITGBASH006434"
        assert order["family"] == "authorization"
        assert order["create_result"] is None
        assert order["order_data"]["type"] == 2
        assert order["order_data"]["orderInfos"][0]["unitPrice"] == "0"

    def test_preview_no_sk_required(self, auth_bytes):
        """preview 模式不要求 sk 头（零下游调用）。"""
        client = TestClient(app)
        assert _post(client, auth_bytes).status_code == 200

    def test_unknown_family_400(self, unknown_bytes):
        client = TestClient(app)
        r = _post(client, unknown_bytes)
        assert r.status_code == 400
        body = r.json()
        # 统一外壳错误：code 机器可读、msg 中文可展示、data 为详情
        assert body["code"] == "unknown_manifest_family"
        assert body["msg"] == "舱单未识别，请使用支持的舱单"
        assert "error" not in body

    def test_preview_box_rejected_200_msg_detail(self, auth_bytes):
        """preview 整批被拒（箱型白名单外）：外壳 code="200"、msg 给具体原因。"""
        wb = load_workbook(io.BytesIO(auth_bytes))
        wb.active.cell(15, 8).value = "1*40GOH (FFAU7731669/SITR853037)"
        buf = io.BytesIO()
        wb.save(buf)
        client = TestClient(app)
        r = _post(client, buf.getvalue())
        assert r.status_code == 200
        body = r.json()
        assert body["code"] == "200"
        assert "40GOH" in body["msg"]
        data = body["data"]
        assert data["summary"] is None  # preview：summary 仍为 null
        assert data["orders"][0]["create_result"]["error"]["code"] == "unknown_box_type"


class TestRouteCreate:
    """create：缺 sk 400；成功 200 + summary；重复上传照常 200。"""

    def test_create_missing_sk_400(self, auth_bytes):
        client = TestClient(app)
        r = _post(client, auth_bytes, data={"create_order": "true"})
        assert r.status_code == 400
        body = r.json()
        assert body["code"] == "bad_request"
        assert "sk" in body["msg"]
        assert "error" not in body

    def test_create_success(self, auth_bytes, monkeypatch):
        async def fake(_payload, _sk):
            return {"success": True, "sn": "11801", "error": None, "upstream": {"bId": 11801}}

        monkeypatch.setattr(service_module, "submit_manifest_async", fake)
        client = TestClient(app)
        r = _post(
            client,
            auth_bytes,
            data={"create_order": "true"},
            headers=CREATE_HEADERS,
        )
        assert r.status_code == 200
        body = r.json()
        assert body["code"] == "200"
        assert body["msg"] == "添加成功"
        data = body["data"]
        assert data["summary"]["created"] == 1
        assert data["summary"]["success_sns"] == ["11801"]
        assert data["upstream"] == {"code": "200", "msg": "成功", "data": [{"bId": 11801}]}

    def test_duplicate_upload_returns_200(self, auth_bytes, monkeypatch):
        """v1.9 放开本地去重：同一文件重复上传照常重新创建，不再 409。"""
        async def fake(_payload, _sk):
            return {"success": True, "sn": "11801", "error": None, "upstream": {"bId": 11801}}

        monkeypatch.setattr(service_module, "submit_manifest_async", fake)
        client = TestClient(app)
        assert _post(
            client, auth_bytes, data={"create_order": "true"}, headers=CREATE_HEADERS
        ).status_code == 200
        r = _post(client, auth_bytes, data={"create_order": "true"}, headers=CREATE_HEADERS)
        assert r.status_code == 200
        body = r.json()
        assert body["code"] == "200"
        assert body["msg"] == "添加成功"
        data = body["data"]
        assert data["summary"]["created"] == 1
        assert data["summary"]["skipped"] == 0
        assert data["upstream"] == {"code": "200", "msg": "成功", "data": [{"bId": 11801}]}

    def test_create_all_failed_204_upstream(self, auth_bytes, monkeypatch):
        """下游全拒 → 200 + 外壳 204 + 具体失败原因（非 HTTP 错误）。"""
        async def fake(_payload, _sk):
            return {"success": False, "sn": None, "error": {"code": "x", "message": "rejected"}}

        monkeypatch.setattr(service_module, "submit_manifest_async", fake)
        client = TestClient(app)
        r = _post(client, auth_bytes, data={"create_order": "true"}, headers=CREATE_HEADERS)
        assert r.status_code == 200
        body = r.json()
        # 外壳业务码 204，msg 透传失败原因（create_result.error.message）
        assert body["code"] == "204"
        assert body["msg"] == "rejected"
        data = body["data"]
        assert data["summary"]["failed"] == 1
        assert data["upstream"] == {"code": "204", "msg": "添加失败", "data": []}
