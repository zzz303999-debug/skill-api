"""bill 解析公开入口与格式分派（P7-1 自 parser.py 拆分）：

- parse_bill：同步入口（测试/纯 CPU 语义基准，内联 L1/L2 与精确回退）；
- open_and_identify / parse_ai_header / parse_exact_fallback：两段式
  编排信号（编排层 to_thread 执行，AiHeaderNeeded 跨段传递 view）；
- _detect_format* / _open_view / _parse_with_engine：xls/xlsx 引擎分派。
依赖方向：opener → parser（解析核心）/ sheet_view，单向无环。"""

from __future__ import annotations

import zipfile
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import BadRequestError, ConvertError

from ..schema import (
    BillPeriod,
    BillRow,
)
from .ai_header import format_header_zone
from .parser import (
    _XLS_MAGIC,
    _XLSX_MAGIC,
    SUPPORTED_EXTS,
    ParseOutput,
    _match_template_and_parse,
    _parse_exact,
    _parse_sheet,
    _parse_with_ai_result,
)
from .sheet_view import _openpyxl_merged_map, _SheetView, _xls_cell_value, _xls_merged_map


def _open_view(path: str | Path, engine: str):
    """打开 workbook 并构造统一视图：返回 (view, 引擎名, 文件名, 资源释放回调)。

    资源释放责任在调用方：同步路径（_parse_with_engine）finally 即时释放；
    两段式（open_and_identify）经 AiHeaderNeeded 移交编排层统一释放。
    文件损坏/无法打开抛 ConvertError（与 parse_bill 同口径）。
    """
    if engine == "xls":
        try:
            book = xlrd.open_workbook(path, formatting_info=True)
        except (xlrd.XLRDError, xlrd.compdoc.CompDocError, OSError) as exc:
            raise ConvertError(
                "failed to open workbook: file is corrupted, encrypted, or not a valid xls",
                details={"path": str(path)},
            ) from exc
        sheet = book.sheet_by_index(0)
        view = _SheetView(
            sheet.nrows,
            sheet.ncols,
            _xls_merged_map(sheet),
            lambda row, col: _xls_cell_value(sheet, row - 1, col - 1),
        )
        return view, "xlrd", Path(path).name, book.release_resources
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise ConvertError(
            "failed to open workbook: file is corrupted, encrypted, or not a valid xlsx",
            details={"path": str(path)},
        ) from exc
    ws = wb.active
    view = _SheetView(
        ws.max_row,
        ws.max_column,
        _openpyxl_merged_map(ws),
        lambda row, col: ws.cell(row, col).value,
    )
    return view, "openpyxl", Path(path).name, wb.close


def _parse_with_engine(path: str | Path, engine: str) -> ParseOutput:
    """按指定引擎打开并解析（同步完整路径）：返回完整 ParseOutput（含未识别列清单）。

    文件损坏/无法打开抛 ConvertError（422 convert_error）；
    结算区间识别不到时返回空 BillPeriod（不报错）。
    """
    view, engine_name, filename, closer = _open_view(path, engine)
    try:
        return _parse_sheet(view, engine_name, filename)
    finally:
        closer()


def parse_xlsx(path: str | Path) -> tuple[list[BillRow], BillPeriod]:
    """读取 .xlsx/.xlsm 竞品账单（openpyxl）：便捷包装，返回 (rows, period)。

    未识别列清单等完整结果请用 parse_bill；找不到合法表头抛 BadRequestError。
    """
    out = _parse_with_engine(path, "xlsx")
    return out.rows, out.period


def parse_xls(path: str | Path) -> tuple[list[BillRow], BillPeriod]:
    """读取 .xls 竞品账单（xlrd）：便捷包装，返回 (rows, period)。

    未识别列清单等完整结果请用 parse_bill；找不到合法表头抛 BadRequestError。
    """
    out = _parse_with_engine(path, "xls")
    return out.rows, out.period


def _detect_format(head: bytes) -> str:
    """按文件头 magic bytes 识别内容格式：'xlsx' / 'xls'；无法识别抛 ConvertError。"""
    if head.startswith(_XLSX_MAGIC):
        return "xlsx"
    if head.startswith(_XLS_MAGIC):
        return "xls"
    raise ConvertError(
        "cannot recognize file content: not a valid xlsx/xls file",
        details={"magic": head[:8].hex()},
    )


def _detect_format_checked(path: str | Path) -> str:
    """扩展名与内容格式校验（前置闸门）：返回内容格式（'xlsx' / 'xls'）。

    扩展名不受支持 → BadRequestError（400 bad_request）；扩展名与内容格式
    不符 → BadRequestError code=file_format_mismatch（400）。同步 parse_bill
    与两段式 open_and_identify 共用，保证两入口前置校验一致。
    """
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise BadRequestError(
            f"unsupported extension: {ext}",
            details={"supported": list(SUPPORTED_EXTS)},
        )
    with open(path, "rb") as f:
        head = f.read(8)
    engine = _detect_format(head)
    if SUPPORTED_EXTS[ext] != engine:
        raise BadRequestError(
            "file extension does not match file content",
            code="file_format_mismatch",
            details={"extension": ext, "detected_format": engine},
        )
    return engine


def parse_bill(path: str | Path) -> ParseOutput:
    """按内容格式分发解析（同步完整路径）：返回 ParseOutput（rows / period / engine / unmatched_headers）。

    生产两段式编排走 open_and_identify 阶段 API（LLM 网络段真异步）；
    本函数为纯 CPU 语义基准（tests 直调依赖），不含 LLM。
    """
    return _parse_with_engine(path, _detect_format_checked(path))


class AiHeaderNeeded:
    """L3 两段式信号：open_and_identify 未命中模板库时返回，待编排层
    完成 LLM 表头映射后二次进段。

    持有跨段资源（view 与 workbook 释放回调；同一时间仅单线程使用 view，
    无并发访问）；资源责任由第一段移交给调用方，完成后须调 close()。
    """

    __slots__ = ("view", "engine", "filename", "zone_lines", "_closer", "_closed")

    def __init__(self, view, engine: str, filename: str, zone_lines, closer) -> None:
        self.view = view
        self.engine = engine
        self.filename = filename
        self.zone_lines = zone_lines
        self._closer = closer
        self._closed = False

    def close(self) -> None:
        """释放 workbook 资源（幂等）。"""
        if not self._closed:
            self._closed = True
            self._closer()


def open_and_identify(path: str | Path) -> ParseOutput | AiHeaderNeeded:
    """阶段 1（同步 CPU，编排层 to_thread 执行）：打开 workbook → L1/L2 指纹识别解析。

    命中模板库 → ParseOutput（workbook 资源已释放，与 parse_bill 同口径，
    常见路径零额外开销）；未命中 → AiHeaderNeeded（view 跨段传递，资源
    责任移交调用方）。扩展名/内容格式前置校验与 parse_bill 完全一致。
    """
    engine = _detect_format_checked(path)
    view, engine_name, filename, closer = _open_view(path, engine)
    out = _match_template_and_parse(view, engine_name, filename)
    if out is not None:
        closer()
        return out
    return AiHeaderNeeded(
        view=view,
        engine=engine_name,
        filename=filename,
        zone_lines=format_header_zone(view),
        closer=closer,
    )


def parse_ai_header(needed: AiHeaderNeeded, ai) -> ParseOutput:
    """阶段 2（同步 CPU，编排层 to_thread 执行）：AI 映射结果 → 候选模板
    → 配置驱动解析（与同步路径共用 _parse_with_ai_result）。

    workbook 资源不在本函数释放（异常路径也要释放）——由编排层 finally
    统一 close。"""
    return _parse_with_ai_result(needed.view, needed.engine, needed.filename, ai)


def parse_exact_fallback(needed: AiHeaderNeeded) -> ParseOutput:
    """LLM 不可用/响应非法回退（同步 CPU，编排层 to_thread 执行）：
    精确匹配（找不到表头照旧 400），与同步路径回退分支同实现。"""
    return _parse_exact(needed.view, needed.engine)
