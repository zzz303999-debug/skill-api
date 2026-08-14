"""竞品账单 Excel 文件解析：表头识别、数据行读取、结算区间识别。

支持 .xlsx（openpyxl，data_only 读取缓存值）与 .xls（xlrd），通过
_SheetView 统一两个引擎的单元格访问（合并单元格填充、值保真语义一致），
按文件头 magic bytes 分发（parse_bill）。

- 表头行：前 MAX_HEADER_SCAN_ROWS 行内定位；流程分三级：
  1. 模板指纹命中（builtin 内置模板 / 已固化 AI 模板）→ 直接用模板列映射，不调 AI
  2. 未命中 → AI 表头映射（异构/变体模板；校验闸门不过抛 400 header_mapping_rejected）
  3. LLM 不可用/响应非法 → 回退现有精确匹配（找不到表头照旧 400）
- 数据行：表头下方所有非空行 → list[BillRow]，单元格值保真转换
  （数字保留浮点尾巴、日期转完整时间串，清洗留给归集步骤）。
- 结算区间：抬头区含「结算」的文本中提取两个 YYYY-MM-DD；
  识别不到返回空 BillPeriod（不报错）。
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.errors import BadRequestError, ConvertError, LLMError, ParseError

from . import template_store
from .fee_map import canonicalize_fee, is_new_fee_schema
from .normalizers import NORMALIZER_REGISTRY, parse_year_hint, to_number
from .schema import (
    HEADER_ALIASES,
    HEADER_COLUMN_MAP,
    IGNORED_HEADERS,
    RECEIVABLE_FEE_COLUMNS,
    REQUIRED_HEADERS,
    BillPeriod,
    BillRow,
)
from .template import compute_fingerprint, load_template

# 表头行最大扫描行数（抬头区通常 1~5 行）
MAX_HEADER_SCAN_ROWS = 15

# 支持的扩展名 → 期望的内容格式（与 _detect_format 的返回值对应）
SUPPORTED_EXTS: dict[str, str] = {".xls": "xls", ".xlsx": "xlsx", ".xlsm": "xlsx"}

# 结算区间日期，如「结算日期：2018-01-01-2018-12-31」
_PERIOD_DATE_RE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})")

# 表头文本中的空白（含全角空格），真实账单表头存在内部空格（如「客户联系 人」）
_HEADER_WHITESPACE_RE = re.compile(r"[\s\u3000]+")

# 内容格式 magic bytes：xlsx 为 zip（PK\x03\x04），xls 为 OLE2（D0 CF 11 E0）
_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0"


class _SheetView:
    """统一工作表视图：1-based 行列访问，合并单元格按需填充左上角值。

    表头/抬头区用 merged_cell（合并值填充），数据区用 cell（原生值，
    合并区非左上角视为空——真实账单尾部「合计:」行为整行合并）。
    """

    def __init__(
        self,
        nrows: int,
        ncols: int,
        merged_map: dict[tuple[int, int], object],
        get_cell,
    ):
        self.nrows = nrows
        self.ncols = ncols
        self._merged = merged_map
        self._get_cell = get_cell

    def cell(self, row: int, col: int) -> object:
        """原生单元格值（不填充合并区）。"""
        return self._get_cell(row, col)

    def merged_cell(self, row: int, col: int) -> object:
        """单元格值：合并区域内的非左上角坐标返回左上角值。"""
        if (row, col) in self._merged:
            return self._merged[(row, col)]
        return self._get_cell(row, col)


@dataclass
class ParseOutput:
    """解析结果：数据行、结算区间、引擎、未识别列告警清单、模板信息。

    unmatched_headers：表头列未被映射（含 IGNORED_HEADERS 之外）且数据区
    至少有一个非空单元格的列表头原文（按列从左到右），供调用方显式处理，
    不静默丢弃；全空列（装饰列/间距列）不告警。AI 识别到的白名单外费用
    以降级 ignore 列形式并入（"new_fee:费用名"）。
    template：本次使用的模板 {fingerprint, source, name}；回退精确匹配时为 None。
    new_template：AI 新映射生成的待固化模板（对账 matched 后由服务层保存）。
    template_match：template_store 识别命中结果（YAML 模板库 L1/L2），含配置；
    未命中（AI 映射/回退精确匹配）时为 None。
    canonical_rows：模板驱动解析且列为标准字段（CanonicalOrder 字段名）时的
    行字典列表（目标字段 → 原文/归一化值），供新归集管线使用；既有流程语义
    （BillRow 字段模板）为 None，行在 rows（list[BillRow]）。
    """

    rows: list[BillRow]
    period: BillPeriod
    engine: str
    unmatched_headers: list[str] = field(default_factory=list)
    template: dict | None = None
    new_fees: list[str] = field(default_factory=list)
    new_template: Any | None = None
    template_match: Any | None = None
    canonical_rows: list[dict] | None = None


def _normalize_header(text: str) -> str:
    """表头文本归一化：去除全部空白后参与列名匹配（仅表头识别用，数据行不处理）。"""
    return _HEADER_WHITESPACE_RE.sub("", text)


def _column_lookup() -> tuple[dict[str, str], dict[str, str]]:
    """归一化表头 → (BillRow 字段名, 费用标准名) 两个查找表（含列名变体）。"""
    data = dict(HEADER_COLUMN_MAP)
    fees = {name: name for name in RECEIVABLE_FEE_COLUMNS}
    fees.update(HEADER_ALIASES)
    return data, fees


def _header_text(value: object) -> str:
    """表头单元格文本：去首尾空白后参与列名匹配（仅表头识别用，数据行不 strip）。"""
    return str(value).strip() if value is not None else ""


def _openpyxl_merged_map(ws: Worksheet) -> dict[tuple[int, int], object]:
    """openpyxl 合并单元格：左上角值映射到区域内所有坐标（1-based）。"""
    values: dict[tuple[int, int], object] = {}
    for merged in ws.merged_cells.ranges:
        top_left = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                values[(row, col)] = top_left
    return values


def _xls_merged_map(sheet: xlrd.sheet.Sheet) -> dict[tuple[int, int], object]:
    """xlrd 合并单元格：merged_cells 为 (rlo, rhi, clo, chi) 0-based 排他边界。"""
    values: dict[tuple[int, int], object] = {}
    for rlo, rhi, clo, chi in sheet.merged_cells:
        top_left = _xls_cell_value(sheet, rlo, clo)
        for row in range(rlo + 1, rhi + 1):
            for col in range(clo + 1, chi + 1):
                values[(row, col)] = top_left
    return values


def _xls_cell_value(sheet: xlrd.sheet.Sheet, row0: int, col0: int) -> object:
    """xlrd 单元格 → 与 openpyxl 等价的 python 对象：number→float（保留浮点尾巴）、
    date→datetime、text→str、其余（空/错误/布尔）→ None。"""
    ctype = sheet.cell_type(row0, col0)
    value = sheet.cell_value(row0, col0)
    if ctype == xlrd.XL_CELL_NUMBER:
        return float(value)
    if ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(value, sheet.book.datemode)
    if ctype == xlrd.XL_CELL_TEXT:
        return value
    return None


def find_header_row(view: _SheetView) -> int:
    """定位同时含「客户编号」「提单号」的表头行，找不到抛 BadRequestError。"""
    seen: set[str] = set()
    scan_rows = min(view.nrows, MAX_HEADER_SCAN_ROWS)
    for row in range(1, scan_rows + 1):
        headers = [
            _normalize_header(_header_text(view.merged_cell(row, col)))
            for col in range(1, view.ncols + 1)
        ]
        seen.update(headers)
        if all(req in headers for req in REQUIRED_HEADERS):
            return row
    missing = [req for req in REQUIRED_HEADERS if req not in seen]
    raise BadRequestError(
        "no valid bill header row found within the first rows",
        details={
            "required_headers": list(REQUIRED_HEADERS),
            "missing_headers": missing,
            "scanned_rows": scan_rows,
        },
    )


def _to_money(value: object) -> float | None:
    """费用单元格转金额（T9 口径）：数字直转；文本按行拆分逐段解析累加
    （合并单元格多值如 '1300\r\n2200' = 两柜各自费用）；复用 to_number 支持
    千分位/货币符号/负号；任一段解析失败（含中文大写）→ None（该格记 warning）。"""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    parts = [p.strip() for p in re.split(r"[\r\n]+", text) if p.strip()]
    if len(parts) == 1:
        return to_number(parts[0])
    total = 0.0
    for part in parts:
        money = to_number(part)
        if money is None:
            return None
        total += money
    return total


def _header_columns(
    view: _SheetView,
    header_row: int,
    data_lookup: dict[str, str] | None = None,
    fee_lookup: dict[str, str] | None = None,
    ignored_cols: set[int] | None = None,
) -> tuple[dict[int, str], dict[int, str], dict[int, str]]:
    """表头行 → (数据列映射, 费用列映射, 潜在未匹配列 {列号: 原始表头文本})。

    未匹配 = 不在 lookup（默认 HEADER_COLUMN_MAP / 费用列 / 别名表；可注入
    模板/AI 映射），且不在 IGNORED_HEADERS（「当前状态」「已收/付金额」等
    有意忽略列不告警）；AI 明确 ignore 的列（ignored_cols）同样不告警。
    是否上报还取决于数据区是否有非空单元格（全空装饰列不报），见 read_data_rows。
    """
    if data_lookup is None or fee_lookup is None:
        data_lookup, fee_lookup = _column_lookup()
    data_cols: dict[int, str] = {}
    fee_cols: dict[int, str] = {}
    unmatched: dict[int, str] = {}
    for col in range(1, view.ncols + 1):
        raw_text = _header_text(view.merged_cell(header_row, col))
        text = _normalize_header(raw_text)
        if text in data_lookup:
            data_cols[col] = data_lookup[text]
        elif text in fee_lookup:
            fee_cols[col] = fee_lookup[text]
        elif raw_text and text not in IGNORED_HEADERS and col not in (ignored_cols or set()):
            unmatched[col] = raw_text
    return data_cols, fee_cols, unmatched


def read_data_rows(
    view: _SheetView,
    header_row: int,
    *,
    data_lookup: dict[str, str] | None = None,
    fee_lookup: dict[str, str] | None = None,
    ignored_cols: set[int] | None = None,
) -> tuple[list[BillRow], list[str]]:
    """读取表头下方所有非空数据行 → (list[BillRow], 未识别列清单)。

    表头行按归一化列名匹配（默认 HEADER_COLUMN_MAP + 费用列 + 别名；可注入
    模板/AI 列名映射）；数据区直接读单元格原生值，合并单元格只有左上角有值
    （真实账单尾部「合计:」行为整行合并，其余字段应保持空，不做左上角值填充）。
    未识别列仅在其数据区至少有一个非空单元格时上报（全空装饰列/间距列不报；
    AI 明确 ignore 的列不报）。
    """
    data_cols, fee_cols, unmatched_cols = _header_columns(
        view, header_row, data_lookup, fee_lookup, ignored_cols
    )
    unmatched_counts = {col: 0 for col in unmatched_cols}

    rows: list[BillRow] = []
    for row in range(header_row + 1, view.nrows + 1):
        kwargs: dict[str, object] = {}
        fees: dict[str, float] = {}
        has_value = False
        for col in range(1, view.ncols + 1):
            # 数据区不填充合并单元格：合并区域只有左上角有值，其余视为空
            raw = view.cell(row, col)
            if raw is None:
                continue
            text = raw if isinstance(raw, str) else str(raw)
            if not text.strip():
                continue
            has_value = True
            if col in data_cols:
                kwargs[data_cols[col]] = text
            elif col in fee_cols:
                money = _to_money(raw)
                if money is not None:
                    fees[fee_cols[col]] = money
                else:
                    # 非数字费用原文保留（如「合计大写」行中文大写金额），
                    # 供归集对账提取第二锚点；累加侧按数字过滤
                    fees[fee_cols[col]] = text
            elif col in unmatched_counts:
                unmatched_counts[col] += 1
        if has_value:
            kwargs["fees"] = fees
            rows.append(BillRow(**kwargs))
    unmatched = [unmatched_cols[col] for col in unmatched_cols if unmatched_counts[col] > 0]
    return rows, unmatched


def _format_period_date(parts: tuple[str, str, str]) -> str | None:
    """YYYY-M-D 三元组 → YYYY-MM-DD；月份/日期越界返回 None。"""
    year, month, day = (int(p) for p in parts)
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    return f"{year:04d}-{month:02d}-{day:02d}"


def extract_bill_period(view: _SheetView, header_row: int) -> BillPeriod:
    """从表头行之前的抬头区识别结算区间；识别不到返回空 BillPeriod（不报错）。"""
    for row in range(1, header_row):
        for col in range(1, view.ncols + 1):
            raw = view.merged_cell(row, col)
            if raw is None:
                continue
            text = str(raw)
            if "结算" not in text:
                continue
            matches = _PERIOD_DATE_RE.findall(text)
            if len(matches) < 2:
                continue
            start = _format_period_date(matches[0])
            end = _format_period_date(matches[1])
            if start and end:
                return BillPeriod(start=start, end=end)
    return BillPeriod()


# 列名消歧后缀：配置源列名「车牌号#1」→ 参与匹配前取「车牌号」
_COL_ORDINAL_RE = re.compile(r"^(.*?)#(\d+)$")
# 双行表头列名：配置源列名「区块.列名」拆区块与列名
_SECTION_COL_RE = re.compile(r"^(.*?)\.(.+)$")
# 序号列数值判定（row_filter: seq_numeric 口径）
_SEQ_NUMERIC_RE = re.compile(r"^\d+(\.\d+)?$")
# 合计分界列（fee_boundary: total_columns 模式）：业务列右边界 = 首个合计列
_TOTAL_COLUMNS = ("应收合计", "应付合计", "成本合计")
# 双行表头家族的费用区块名（区块行非空单元格即区块名，费用区不参与业务列匹配）
_SECTION_NAMES = ("应收", "应付", "车辆成本", "公司成本", "成本")

# 对账锚点列关键字（T9）：含这些关键字的列不是费目（合计/小计/已收/未收/已付/未付/
# 利润），值保留为对账锚点；bill_total 取「合计/小计」列，其余仅排除不抽取
_ANCHOR_KEYWORDS = ("合计", "小计", "已收", "已付", "未收", "未付", "利润")
# 备注类列名（column_range 家族费用列发现时排除，避免备注被当费用列）
_NOTE_KEYWORDS = ("备注", "附言")


# BillRow 字段名集合：模板 columns 目标字段 ⊆ 该集合 → 既有流程语义（构造 BillRow）
_BILLROW_FIELDS = set(BillRow.model_fields)


def _settlement_text(view: _SheetView, header_row: int) -> str | None:
    """表头上方抬头区含「结算」的文本（year_hint 解析用）。"""
    parts: list[str] = []
    for row in range(1, header_row):
        for col in range(1, view.ncols + 1):
            raw = view.merged_cell(row, col)
            if raw is None:
                continue
            text = str(raw)
            if "结算" in text:
                parts.append(text)
    return "\n".join(parts) if parts else None


def _resolve_year_hint(
    template: dict, filename: str, view: _SheetView, header_row: int
) -> tuple[int, int] | None:
    """按配置 year_source 顺序尝试年份提示：filename → 结算日期行。"""
    sources = template.get("year_source") or []
    settlement_text = _settlement_text(view, header_row)
    for source in sources:
        if source == "settlement_row":
            hint = parse_year_hint("", settlement_text)
        else:
            hint = parse_year_hint(filename or "", None)
        if hint is not None:
            return hint
    return None


def _apply_normalizer(
    field: str, func_name: str, value, year_hint: tuple[int, int] | None
):
    """按配置引用应用归一化函数；返回归一化值（解析失败返回 None，缺失语义）。

    特判：vessel_voyage_split 拆出的二元组按目标字段取对应项（vessel 取船名、
    voyage 取航次）；date_flex 注入 year_hint（M-D 补年份）。
    """
    func = NORMALIZER_REGISTRY.get(func_name)
    if func is None:
        return value
    if func_name == "vessel_voyage_split":
        vessel, voyage = func(value)
        return vessel if field == "vessel" else voyage
    if func_name == "date_flex":
        return func(value, year_hint)
    return func(value)


def _header_column_index(
    view: _SheetView, header_row: int, two_row: bool
) -> tuple[list[str], list[str]]:
    """表头行列名索引 → (列名(归一化), 区块名(归一化，单行全空))。

    双行表头：header_row-1 为区块行（section_fill: forward 即 merged_cell
    合并右填充），列名 = 区块非空时「区块.列名」；单行表头区块全空。
    """
    section_row = header_row - 1 if two_row and header_row > 1 else 0
    sections: list[str] = []
    names: list[str] = []
    for col in range(1, view.ncols + 1):
        section = (
            str(view.merged_cell(section_row, col)).strip() if section_row else ""
        )
        name = str(view.merged_cell(header_row, col)).strip()
        sections.append(_normalize_header(section))
        if section:
            names.append(f"{_normalize_header(section)}.{_normalize_header(name)}")
        else:
            names.append(_normalize_header(name))
    return names, sections


def _business_end_col(
    template: dict, names: list[str], sections: list[str]
) -> int:
    """业务列右边界（fee_boundary 三种模式；本轮只圈定业务列范围，费用值不映射）。

    - section_header：费用区块（应收/应付/车辆成本/公司成本）起始列
    - total_columns：首个合计列（应收合计/应付合计/成本合计）
    - column_range/缺省：无分界，全部列（配置精确控制）
    """
    mode = template.get("fee_boundary", "column_range")
    if mode == "section_header":
        for col, section in enumerate(sections, start=1):
            if section in _SECTION_NAMES:
                return col
        return len(names) + 1
    if mode == "total_columns":
        for col, name in enumerate(names, start=1):
            if any(total in name for total in _TOTAL_COLUMNS):
                return col
        return len(names) + 1
    return len(names) + 1


def _match_source_cols(
    spec, names: list[str], sections: list[str], business_end: int
) -> list[int]:
    """配置源列名规范 → 匹配到的列号列表（业务区内按列序）。

    支持：
    - 「列名」：归一化列名相等（单行家族）；
    - 「区块.列名」：区块与列名双匹配（two_row 家族费用列等带前缀列）；
    - 「列名#N」：同名第 N 次出现（同名列消歧，如 车牌号#1 取首次）。
    """
    if isinstance(spec, list):
        cols: list[int] = []
        for item in spec:
            cols.extend(_match_source_cols(item, names, sections, business_end))
        return sorted(set(cols))
    spec = str(spec).strip()
    if not spec:
        return []
    m = _COL_ORDINAL_RE.match(spec)
    if m:
        col_name, ordinal = m.group(1), int(m.group(2))
        matched = [
            col
            for col, name in enumerate(names, start=1)
            if col < business_end and name == _normalize_header(col_name)
        ]
        return [matched[ordinal - 1]] if ordinal <= len(matched) else []
    m = _SECTION_COL_RE.match(spec)
    if m:
        section, col_name = m.group(1), m.group(2)
        return [
            col
            for col in range(1, len(names) + 1)
            if col < business_end
            and sections[col - 1] == _normalize_header(section)
            and names[col - 1] == _normalize_header(col_name)
        ]
    normalized = _normalize_header(spec)
    return [
        col
        for col, name in enumerate(names, start=1)
        if col < business_end and name == normalized
    ]


def _is_anchor_column(name: str) -> bool:
    """对账锚点列判定：列名含 合计/小计/已收/未收/已付/未付/利润 关键字。"""
    return any(kw in name for kw in _ANCHOR_KEYWORDS)


def _is_note_column(name: str) -> bool:
    """备注类列判定（column_range 家族费用列发现时排除）。"""
    return any(kw in name for kw in _NOTE_KEYWORDS)


def _range_bounds(
    ranges_cfg: dict, names: list[str], field_cols: dict[str, list[int]]
) -> dict[str, tuple[int, int]]:
    """单行 total_columns 家族：区块 → (费用区起点列, 终点列)（排他）。

    区块费用区 = 上一分界列之后 到 本区块分界列（ranges 值 = 区块末尾合计列名）
    之前；首个区块起点 = 业务列（columns 匹配）最后一列之后。分界列缺失 → 区块跳过。
    """
    business_last = max((c for cols in field_cols.values() for c in cols), default=0)
    bounds: dict[str, tuple[int, int]] = {}
    prev_end = business_last + 1
    for section, anchor_name in ranges_cfg.items():
        end = next(
            (
                col
                for col, name in enumerate(names, start=1)
                if name == _normalize_header(str(anchor_name))
            ),
            None,
        )
        if end is None:
            continue
        bounds[section] = (prev_end, end)
        prev_end = end + 1
    return bounds


def _discover_fee_columns(
    fees_cfg: dict,
    names: list[str],
    sections: list[str],
    field_cols: dict[str, list[int]],
    row_anchor: str = "序号",
) -> tuple[dict[int, tuple[str, str]], dict[int, tuple[str, str]]]:
    """新 fees schema（T9）：费用列/锚点列发现，代码零费目名硬编码。

    返回 (费目列 {列号: (区块, 列名)}, 锚点列 {列号: (区块, 列名)})：
    - two_row（区块行非空）：区块 ∈ fees.channels 的列；锚点列 = 含锚点关键字列；
    - 单行 + fees.ranges（区块 → 末尾合计列名）：范围内列（长尾费目也抽取 → to_other），
      边界合计列本身是锚点列（对账锚点）；
    - 单行无 ranges（column_range 家族）：columns 未映射列中排除 row_anchor（序号）/
      备注/IGNORED 后的全部列（mapping 只负责码解析，长尾自动 to_other，保证对账恒等）；
      锚点列取 fees.anchors 显式声明（账单侧合计列语义家族差异大，配置为准）。
    """
    fee_cols: dict[int, tuple[str, str]] = {}
    anchor_cols: dict[int, tuple[str, str]] = {}
    anchor_row = _normalize_header(row_anchor)
    if any(sections):
        channels = set(fees_cfg.get("channels") or {})
        for col, (section, name) in enumerate(zip(sections, names, strict=True), start=1):
            if not section or section not in channels or not name:
                continue
            # two_row 表头列名为「区块.列名」复合键，抽取侧还原纯列名
            # （区块参与匹配由 section 承担，费目名/字典/note 都须无前缀）
            if name.startswith(section + "."):
                name = name[len(section) + 1 :]
            if _is_note_column(name):
                continue  # 备注类列（防御：区块内误排备注）
            if _is_anchor_column(name):
                anchor_cols[col] = (section, name)
            else:
                fee_cols[col] = (section, name)
        return fee_cols, anchor_cols
    ranges = fees_cfg.get("ranges") or {}
    if ranges:
        bounds = _range_bounds(ranges, names, field_cols)
        boundary_section = {_normalize_header(str(v)): s for s, v in ranges.items()}
        for col, name in enumerate(names, start=1):
            if not name:
                continue
            section = next(
                (s for s, (lo, hi) in bounds.items() if lo <= col < hi), None
            )
            if section is None:
                # 边界合计列（ranges 值）本身是锚点，归属声明区块
                section = boundary_section.get(_normalize_header(name))
                if section is not None:
                    anchor_cols[col] = (section, name)
                continue
            if _is_note_column(name):
                continue  # 备注类列（如赢辉大表备注，非费目）
            if _is_anchor_column(name):
                anchor_cols[col] = (section, name)
            else:
                fee_cols[col] = (section, name)
        return fee_cols, anchor_cols
    # column_range：columns 未映射列（排除序号/锚点/备注/IGNORED）即费用列；
    # 锚点列取 fees.anchors 显式声明（如 秋怡小计列；军羽小计列恒值无效则不配）
    mapped = {c for cols in field_cols.values() for c in cols}
    for col, name in enumerate(names, start=1):
        if col in mapped or not name:
            continue
        if _normalize_header(name) == anchor_row:
            continue  # 序号列（row_anchor）非费用
        if _is_anchor_column(name):
            continue  # 合计/小计等锚点列非费目（锚点只取 fees.anchors 显式声明）
        if _is_note_column(name) or name in IGNORED_HEADERS:
            continue
        fee_cols[col] = ("", name)
    anchors_cfg = fees_cfg.get("anchors") or {}
    for section, anchor_names in anchors_cfg.items():
        for anchor_name in anchor_names if isinstance(anchor_names, list) else [anchor_names]:
            for col, name in enumerate(names, start=1):
                if name == _normalize_header(str(anchor_name)):
                    anchor_cols[col] = (str(section), name)
                    fee_cols.pop(col, None)
    return fee_cols, anchor_cols


def _parse_with_template(
    view: _SheetView, match, engine: str, filename: str
) -> ParseOutput:
    """配置驱动解析（L1/L2 命中后）：表头定位/区块/列映射/行过滤/归一化全由
    模板配置声明，代码对目标字段名零假设。

    输出分流（按模板 columns 目标字段集合判定）：
    - 目标 ⊆ BillRow 字段（内置模板迁移语义）→ rows: list[BillRow]（原文保真，
      费用进 fees；清洗/日期补年仍由归集层负责，与迁移前一致）；
    - 标准字段（CanonicalOrder 字段名）→ canonical_rows: list[dict]（归一化后）。
    """
    template = match.template
    header_row = match.header_row
    header_cfg = template.get("header", {}) or {}
    two_row = bool(header_cfg.get("two_row"))
    names, sections = _header_column_index(view, header_row, two_row)
    business_end = _business_end_col(template, names, sections)

    # columns 配置 → 目标字段 → 源列号（保持配置书写顺序）
    columns = template.get("columns", {}) or {}
    field_cols: dict[str, list[int]] = {}
    for target_field, spec in columns.items():
        cols = _match_source_cols(spec, names, sections, business_end)
        if cols:
            field_cols[target_field] = cols

    # fees 段：新 schema（T10 费用通道配置）→ 费用列/锚点列发现 + 行级费用抽取；
    # 旧 schema（内置模板迁移：费用名 → 源列名）→ BillRow.fees 抽取（语义不变）
    fees_cfg = template.get("fees", {}) or {}
    fee_cols: dict[int, str] = {}
    new_fee_cols: dict[int, tuple[str, str]] = {}
    anchor_cols: dict[int, tuple[str, str]] = {}
    channels: dict = {}
    if is_new_fee_schema(fees_cfg):
        row_anchor = _normalize_header(str(header_cfg.get("row_anchor") or "序号"))
        new_fee_cols, anchor_cols = _discover_fee_columns(
            fees_cfg, names, sections, field_cols, row_anchor
        )
        channels = fees_cfg.get("channels") or {}
    else:
        for fee_name, spec in fees_cfg.items():
            for col in _match_source_cols(spec, names, sections, business_end):
                fee_cols[col] = fee_name

    # 未识别列（非 IGNORED、非映射、数据区非空才上报）；unmatched 存原始表头文本；
    # AI 明确忽略的列（L3 候选配置 _ignored_cols）不告警
    ignored_cols = set(template.get("_ignored_cols") or [])
    mapped_cols_flat = (
        {c for cols in field_cols.values() for c in cols}
        | set(fee_cols)
        | set(new_fee_cols)
        | set(anchor_cols)
    )
    unmatched_raw: dict[int, str] = {
        col: str(view.merged_cell(header_row, col)).strip()
        for col, name in enumerate(names, start=1)
        if (
            col not in mapped_cols_flat
            and col not in ignored_cols
            and name
            and name not in IGNORED_HEADERS
        )
    }
    unmatched_hits: dict[int, int] = {col: 0 for col in unmatched_raw}

    # 序号列（row_filter: seq_numeric 口径）：row_anchor 列，columns 有 seq 时优先
    anchor = _normalize_header(str(header_cfg.get("row_anchor") or "序号"))
    seq_cols = [col for col, name in enumerate(names, start=1) if name == anchor]
    if "seq" in field_cols:
        seq_cols = field_cols["seq"]
    seq_col = seq_cols[0] if seq_cols else None

    stop_on = [str(s) for s in (template.get("data", {}) or {}).get("stop_on") or []]
    row_filter = (template.get("data", {}) or {}).get("row_filter", "seq_numeric")
    normalizers = template.get("normalizers", {}) or {}
    year_hint = _resolve_year_hint(template, filename, view, header_row)

    is_billrow = set(columns) <= _BILLROW_FIELDS
    raw_rows: list[BillRow] = [] if is_billrow else []
    canonical_rows: list[dict] = [] if not is_billrow else []

    for row in range(header_row + 1, view.nrows + 1):
        # stop_on：任一列值含终止词 → 停止（合计/制单人页脚等）
        if stop_on:
            row_texts = [
                str(view.cell(row, col)).strip()
                for col in range(1, view.ncols + 1)
                if view.cell(row, col) is not None
            ]
            if any(any(s in t for s in stop_on) for t in row_texts):
                break
        # 未识别列计数（数据区非空才上报）
        for col in unmatched_hits:
            raw = view.cell(row, col)
            if raw is not None and str(raw).strip():
                unmatched_hits[col] += 1
        # row_filter: seq_numeric——序号列须为数值才是数据行；
        # BillRow 语义模板（内置迁移）不生效：尾部合计/箱量/大写行是归集对账锚点，
        # 须保留（与迁移前「全部非空行」一致，过滤发生在归集步骤）
        if row_filter == "seq_numeric" and seq_col is not None and not is_billrow:
            seq_raw = view.cell(row, seq_col)
            if seq_raw is None or not _SEQ_NUMERIC_RE.match(str(seq_raw).strip()):
                continue

        values: dict[str, object] = {}
        fees_map: dict[str, float | str] = {}
        for target_field, cols in field_cols.items():
            for col in cols:
                raw = view.cell(row, col)
                if raw is None:
                    continue
                text = str(raw).strip()
                if not text:
                    continue
                # 多源列取最后一个非空（保持既有覆盖语义，空值不覆盖）
                values[target_field] = text
        for col, fee_name in fee_cols.items():
            raw = view.cell(row, col)
            if raw is None:
                continue
            if str(raw).strip():
                money = _to_money(raw)
                fees_map[fee_name] = money if money is not None else str(raw)

        # 新 fees schema 行级费用抽取（T9）：费用列 → 行级费用项（随归集按键汇总）；
        # 金额解析失败/空值/0 不生成记录，失败记 warning 进对账报告；锚点列值保留
        fee_items: list[dict] = []
        anchors: dict[str, dict[str, float]] = {}
        fee_failures: list[dict] = []
        fee_skipped: list[dict] = []
        if new_fee_cols:
            # column_range 家族单行表头无区块行，section 默认取唯一启用区块
            # （保证「应收.费目」复合键映射与通道归属正确）
            default_section = next(iter(channels), "")
            for col, (section, name) in new_fee_cols.items():
                section = section or default_section
                raw = view.cell(row, col)
                if raw is None or not str(raw).strip():
                    continue
                money = _to_money(raw)
                if money is None:
                    fee_failures.append({"section": section, "name": name, "raw": str(raw)})
                    continue
                if money == 0:
                    continue  # 空值与 0 均不生成费用记录（账单侧合计仍按列值参与对账）
                meta = canonicalize_fee(section, name, fees_cfg)
                if not meta["code"]:
                    fee_skipped.append({"section": section, "name": name, "money": money})
                    continue
                # T27a 负向扣减项：negative=true → 金额取负录入（两位小数字符串如 "-50.00"）；
                # 账单语义：应付合计 = Σ正项 − 扣除费列值（列值本身带符号：+486 扣减、-25 加回），
                # 故取 `-money`（相反数）而非 -abs：负值列取负后为 +（加回），恒等式自然成立；
                # negative_policy=skip_report 已在 canonicalize_fee 降级 import=False（不录入仅对账）
                if meta.get("negative"):
                    money = -money
                fee_items.append(
                    {
                        "section": section,
                        "name": name,
                        "channel": channels.get(section) or "shou",
                        "code": meta["code"],
                        "import": meta["import"],
                        "reconcile": meta["reconcile"],
                        "negative": bool(meta.get("negative")),
                        "money": money,
                    }
                )
            for col, (section, name) in anchor_cols.items():
                raw = view.cell(row, col)
                if raw is None or not str(raw).strip():
                    continue
                money = _to_money(raw)
                if money is None:
                    continue
                anchors.setdefault(section, {})[name] = money
        if not values and not fees_map and not fee_items:
            if is_billrow:
                # 与迁移前 read_data_rows 一致：仅未映射列有值的行也保留（字段为空）
                has_any = any(
                    view.cell(row, col) is not None and str(view.cell(row, col)).strip()
                    for col in range(1, view.ncols + 1)
                )
                if not has_any:
                    continue
            else:
                continue

        if is_billrow:
            kwargs = dict(values)
            if fees_map:
                kwargs["fees"] = fees_map
            raw_rows.append(BillRow(**kwargs))
        else:
            normalized: dict[str, object] = {}
            for target_field, value in values.items():
                func_name = normalizers.get(target_field)
                if func_name:
                    value = _apply_normalizer(
                        target_field, func_name, value, year_hint
                    )
                if value is not None and value != "":
                    normalized[target_field] = value
            # 行级费用私有键（归集侧按 group_key 聚合；下划线前缀不进响应）
            if fee_items:
                normalized["_fees"] = fee_items
            if anchors:
                normalized["_anchors"] = anchors
            if fee_failures:
                normalized["_fee_failures"] = fee_failures
            if fee_skipped:
                normalized["_fee_skipped"] = fee_skipped
            canonical_rows.append(normalized)

    unmatched = [
        unmatched_raw[col]
        for col in range(1, view.ncols + 1)
        if col in unmatched_hits and unmatched_hits[col] > 0
    ]
    period = extract_bill_period(view, header_row)
    template_meta = {
        "fingerprint": match.fingerprint,
        "source": "builtin" if is_billrow else "template",
        "name": template.get("name", match.template_id),
    }
    return ParseOutput(
        rows=raw_rows,
        period=period,
        engine=engine,
        unmatched_headers=unmatched,
        template=template_meta,
        template_match=match,
        canonical_rows=canonical_rows if not is_billrow else None,
    )


def _parse_sheet(view: _SheetView, engine: str, filename: str = "") -> ParseOutput:
    """统一解析核心：YAML 模板库识别（L1 指纹/L2 族级近似）→ 配置驱动解析；
    未命中 → 旧指纹库/AI 映射（回退精确匹配）。"""
    match = template_store.identify(view)
    if match is not None:
        return _parse_with_template(view, match, engine, filename)
    scan_rows = min(view.nrows, MAX_HEADER_SCAN_ROWS)
    # 1) 指纹命中模板库（builtin 内置模板 / 已固化 AI 模板）→ 直接模板列映射，不调 AI
    for row in range(1, scan_rows + 1):
        headers = [
            _normalize_header(_header_text(view.merged_cell(row, col)))
            for col in range(1, view.ncols + 1)
        ]
        template = load_template(compute_fingerprint(headers))
        if template is None:
            continue
        rows, unmatched = read_data_rows(
            view,
            row,
            data_lookup=dict(template.column_map),
            fee_lookup=dict(template.fee_map),
        )
        period = extract_bill_period(view, row)
        return ParseOutput(
            rows=rows,
            period=period,
            engine=engine,
            unmatched_headers=unmatched,
            template={
                "fingerprint": template.fingerprint,
                "source": template.source,
                "name": template.name,
            },
        )
    # 2) 未命中模板库 → L3 AI 表头映射（标准字段映射 + 模板结构判定；
    #    四道校验闸门不过抛 400；LLM 不可用/响应非法 → 回退现有精确匹配）
    try:
        from .ai_header import map_header

        ai = map_header(view)
    except (LLMError, ParseError):
        # LLM 不可用/响应非法 → 回退现有精确匹配（找不到表头照旧 400）
        return _parse_exact(view, engine)
    headers = [
        _normalize_header(_header_text(view.merged_cell(ai.header_row, col)))
        for col in range(1, view.ncols + 1)
    ]
    fingerprint = template_store.compute_fingerprint(headers)
    # L3 映射 → 临时模板配置（columns 段 + two_row/fee_boundary）→ 复用配置驱动解析：
    # 与 YAML 模板同一条代码路径；费用列本轮不映射（AI 白名单费用降级上报）
    candidate = template_store.build_template_config(
        fingerprint=fingerprint,
        header_row=ai.header_row,
        two_row=ai.two_row,
        fee_boundary=ai.fee_boundary,
        column_map=ai.column_map,
        ignored_cols=ai.ignored_cols,
    )
    match = template_store.TemplateMatch(
        template=candidate,
        header_row=ai.header_row,
        level="L3",
        fingerprint=fingerprint,
        template_id=candidate["template_id"],
    )
    out = _parse_with_template(view, match, engine, filename)
    # 白名单外费用降级上报（new_fee: 前缀，与既有语义一致）
    out.unmatched_headers = [
        *out.unmatched_headers,
        *(f"new_fee:{name}" for name in ai.new_fees),
    ]
    out.new_fees = ai.new_fees
    # 候选模板配置：预览界面展示映射结果，人工确认后固化为 templates/{family}_v1.yaml
    out.new_template = candidate
    return out


def _parse_exact(view: _SheetView, engine: str) -> ParseOutput:
    """现状精确匹配路径（LLM 不可用回退）：找不到表头照旧 400，无模板信息。"""
    header_row = find_header_row(view)
    rows, unmatched = read_data_rows(view, header_row)
    period = extract_bill_period(view, header_row)
    return ParseOutput(rows=rows, period=period, engine=engine, unmatched_headers=unmatched)


def _parse_with_engine(path: str | Path, engine: str) -> ParseOutput:
    """按指定引擎打开并解析：返回完整 ParseOutput（含未识别列清单）。

    文件损坏/无法打开抛 ConvertError（422 convert_error）；
    结算区间识别不到时返回空 BillPeriod（不报错）。
    """
    if engine == "xls":
        try:
            book = xlrd.open_workbook(path, formatting_info=True)
        except (xlrd.XLRDError, xlrd.compdoc.CompDocError, OSError) as exc:
            raise ConvertError(
                "failed to open workbook: file is corrupted, encrypted, or not a valid xls",
                details={"path": str(path)},
            ) from exc
        try:
            sheet = book.sheet_by_index(0)
            view = _SheetView(
                sheet.nrows,
                sheet.ncols,
                _xls_merged_map(sheet),
                lambda row, col: _xls_cell_value(sheet, row - 1, col - 1),
            )
            return _parse_sheet(view, "xlrd", Path(path).name)
        finally:
            book.release_resources()
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise ConvertError(
            "failed to open workbook: file is corrupted, encrypted, or not a valid xlsx",
            details={"path": str(path)},
        ) from exc
    try:
        ws = wb.active
        view = _SheetView(
            ws.max_row,
            ws.max_column,
            _openpyxl_merged_map(ws),
            lambda row, col: ws.cell(row, col).value,
        )
        return _parse_sheet(view, "openpyxl", Path(path).name)
    finally:
        wb.close()


def parse_xlsx(path: str | Path) -> tuple[list[BillRow], BillPeriod]:
    """读取 .xlsx/.xlsm 竞品账单（openpyxl）：便捷包装，返回 (rows, period)。

    未识别列清单等完整结果请用 parse_bill；找不到合法表头抛 BadRequestError。
    """
    out = _parse_with_engine(path, "xlsx")
    return out.rows, out.period


def parse_xls(path: str | Path) -> tuple[list[BillRow], BillPeriod]:
    """读取 .xls 竞品账单（xlrd）：便捷包装，返回 (rows, period)。

    未识别列清单等完整结果请用 parse_bill；找不到合法表头抛 BadRequestError。
    """
    out = _parse_with_engine(path, "xls")
    return out.rows, out.period


def _detect_format(head: bytes) -> str:
    """按文件头 magic bytes 识别内容格式：'xlsx' / 'xls'；无法识别抛 ConvertError。"""
    if head.startswith(_XLSX_MAGIC):
        return "xlsx"
    if head.startswith(_XLS_MAGIC):
        return "xls"
    raise ConvertError(
        "cannot recognize file content: not a valid xlsx/xls file",
        details={"magic": head[:8].hex()},
    )


def parse_bill(path: str | Path) -> ParseOutput:
    """按内容格式分发解析：返回 ParseOutput（rows / period / engine / unmatched_headers）。

    扩展名不受支持 → BadRequestError（400 bad_request）；
    扩展名与内容格式不符 → BadRequestError code=file_format_mismatch（400）。
    """
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise BadRequestError(
            f"unsupported extension: {ext}",
            details={"supported": list(SUPPORTED_EXTS)},
        )
    with open(path, "rb") as f:
        head = f.read(8)
    engine = _detect_format(head)
    if SUPPORTED_EXTS[ext] != engine:
        raise BadRequestError(
            "file extension does not match file content",
            code="file_format_mismatch",
            details={"extension": ext, "detected_format": engine},
        )
    if engine == "xls":
        return _parse_with_engine(path, "xls")
    return _parse_with_engine(path, "xlsx")
