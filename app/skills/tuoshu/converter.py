#!/usr/bin/env python3
"""
tuoshu-extractor: 任意格式 → Markdown（无损转换）

职责单一：xlsx/xls/docx/doc/pdf → stdout markdown。
不做任何语义抽取、不做归一化、不做 OCR。
图片和扫描 PDF 输出 SCAN_OR_IMAGE_HINT: <path>，由上层转为 vision 输入。

用法:
    python3 scripts/to_text.py <file>
"""

from __future__ import annotations

import io
import os
import posixpath
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree

# ---------- 通用工具 ----------

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".gif", ".webp"}


def emit_scan_hint(path: str, reason: str = "") -> None:
    """输出 OCR 提示，供上层 skill 调用 OCR。"""
    suffix = f"  # {reason}" if reason else ""
    print(f"SCAN_OR_IMAGE_HINT: {path}{suffix}")


def _cell_col_letter(idx_1based: int) -> str:
    """1 → A, 27 → AA。"""
    s = ""
    n = idx_1based
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _clean_cell(value) -> str:
    """把任意单元格值转成字符串，保留换行为 <br>。"""
    if value is None:
        return ""
    s = str(value)
    # 统一换行 → <br>，保结构
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = s.replace("\n", "<br>")
    # 去尾部空白，不去内部空白
    return s.rstrip()


def _format_xlsx_value(value, number_format: str) -> str:
    """Format date/time cells without inventing a midnight component."""
    fmt = (number_format or "").lower()
    has_date = any(token in fmt for token in ("yy", "dd"))
    has_time = any(token in fmt for token in ("h", "s"))
    has_seconds = "s" in fmt
    if isinstance(value, datetime):
        if has_date and has_time:
            pattern = "%Y-%m-%d %H:%M:%S" if has_seconds else "%Y-%m-%d %H:%M"
            return value.strftime(pattern)
        if has_time and not has_date:
            return value.strftime("%H:%M:%S" if has_seconds else "%H:%M")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S" if has_seconds else "%H:%M")
    return _clean_cell(value)


# ---------- 分支：xlsx ----------

def convert_xlsx(path: str) -> dict[str, object]:
    import openpyxl

    # 用 BytesIO 绕过 openpyxl 对扩展名的检查（有些 .xls 实际是 xlsx）
    with open(path, "rb") as fh:
        data = fh.read()
    value_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    formula_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    print(f"# {Path(path).name}")
    print("_format: xlsx_")
    print()

    formula_count = 0
    formula_values_available = 0
    formula_values_missing: list[str] = []
    merged_range_count = 0
    nonempty_cell_count = 0

    for ws in value_wb.worksheets:
        formula_ws = formula_wb[ws.title]
        print(f"## Sheet: {ws.title}")
        print()

        # Only the anchor owns a merged value. Repeating it into every covered
        # cell creates false records and makes Markdown diverge from Excel.
        merged_continuations: set[tuple[int, int]] = set()
        merged_ranges: list[str] = []
        for rng in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = rng.min_col, rng.min_row, rng.max_col, rng.max_row
            merged_ranges.append(str(rng))
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    if (r, c) != (min_row, min_col):
                        merged_continuations.add((r, c))
        merged_range_count += len(merged_ranges)

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        if max_row == 0 or max_col == 0:
            print("_empty sheet_")
            print()
            continue

        # 修剪尾部纯空的行/列（结构中间的空行空列保留，只去外围的填充空白）
        def _cell_val(
            r: int,
            c: int,
            worksheet=ws,
            formulas=formula_ws,
            merged=merged_continuations,
        ):
            if (r, c) in merged:
                return None
            value = worksheet.cell(row=r, column=c).value
            if value is not None:
                return value
            return formulas.cell(row=r, column=c).value

        while max_row > 0 and all(
            _cell_val(max_row, c) in (None, "") for c in range(1, max_col + 1)
        ):
            max_row -= 1
        while max_col > 0 and all(
            _cell_val(r, max_col) in (None, "") for r in range(1, max_row + 1)
        ):
            max_col -= 1

        if max_row == 0 or max_col == 0:
            print("_empty sheet_")
            print()
            continue

        # 输出 markdown 表格（保留结构中的空行/空列）
        header = "| " + " | ".join(["_row/col_"] + [_cell_col_letter(c) for c in range(1, max_col + 1)]) + " |"
        sep = "| " + " | ".join(["---"] * (max_col + 1)) + " |"
        print(header)
        print(sep)

        for r in range(1, max_row + 1):
            row_cells: list[str] = [str(r)]
            for c in range(1, max_col + 1):
                if (r, c) in merged_continuations:
                    row_cells.append("")
                    continue
                value_cell = ws.cell(row=r, column=c)
                formula_cell = formula_ws.cell(row=r, column=c)
                value = value_cell.value
                is_formula = isinstance(formula_cell.value, str) and formula_cell.value.startswith("=")
                if value not in (None, "") or is_formula:
                    nonempty_cell_count += 1
                if is_formula:
                    formula_count += 1
                    if value is None:
                        formula_values_missing.append(f"{ws.title}!{value_cell.coordinate}")
                    else:
                        formula_values_available += 1
                rendered = _format_xlsx_value(value, value_cell.number_format)
                row_cells.append(rendered.replace("|", "\\|"))
            print("| " + " | ".join(row_cells) + " |")
        print()

        if merged_ranges:
            print("### Merged ranges")
            print()
            for merged_range in merged_ranges:
                print(f"- `{merged_range}`")
            print()

        formula_rows: list[tuple[str, str, str]] = []
        for row in formula_ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for formula_cell in row:
                formula = formula_cell.value
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                value_cell = ws[formula_cell.coordinate]
                cached = (
                    _format_xlsx_value(value_cell.value, value_cell.number_format)
                    if value_cell.value is not None
                    else "_unavailable_"
                )
                formula_rows.append((formula_cell.coordinate, formula, cached))
        if formula_rows:
            print("### Formula cells")
            print()
            print("| Cell | Formula | Cached value |")
            print("| --- | --- | --- |")
            for coordinate, formula, cached in formula_rows:
                print(
                    "| "
                    + " | ".join(
                        value.replace("|", "\\|") for value in (coordinate, formula, cached)
                    )
                    + " |"
                )
            print()

    report = {
        "coverage": {
            "sheet_count": len(value_wb.worksheets),
            "nonempty_cells": nonempty_cell_count,
            "merged_ranges": merged_range_count,
            "formula_cells": formula_count,
            "formula_values_available": formula_values_available,
            "complete": not formula_values_missing,
            "omissions": formula_values_missing,
        },
        "formula_values_missing": formula_values_missing,
    }
    value_wb.close()
    formula_wb.close()
    return report


# ---------- 分支：xls ----------

def convert_xls(path: str) -> dict[str, object]:
    import xlrd
    from xlrd import xldate

    # 有些 .xls 文件实际是 xlsx（ZIP 头 PK\x03\x04），转发处理
    try:
        with open(path, "rb") as fh:
            magic = fh.read(4)
    except OSError:
        magic = b""
    if magic == b"PK\x03\x04":
        return convert_xlsx(path)

    book = xlrd.open_workbook(path, formatting_info=False)
    datemode = book.datemode
    print(f"# {Path(path).name}")
    print("_format: xls_")
    print()

    for sheet in book.sheets():
        print(f"## Sheet: {sheet.name}")
        print()

        nrows, ncols = sheet.nrows, sheet.ncols
        if nrows == 0 or ncols == 0:
            print("_empty sheet_")
            print()
            continue

        # xlrd 的 merged_cells 是 [(rlo, rhi, clo, chi), ...]，含前不含后
        merged_continuations: set[tuple[int, int]] = set()
        merged_ranges: list[str] = []
        for rlo, rhi, clo, chi in sheet.merged_cells:
            merged_ranges.append(
                f"{_cell_col_letter(clo + 1)}{rlo + 1}:{_cell_col_letter(chi)}{rhi}"
            )
            for r in range(rlo, rhi):
                for c in range(clo, chi):
                    if (r, c) != (rlo, clo):
                        merged_continuations.add((r, c))

        def _xls_val(
            r: int,
            c: int,
            current_sheet=sheet,
            merged=merged_continuations,
        ):
            if (r, c) in merged:
                return None
            ctype = current_sheet.cell_type(r, c)
            v = current_sheet.cell_value(r, c)
            # XL_CELL_DATE = 3：还原为 ISO datetime 字符串
            if ctype == xlrd.XL_CELL_DATE:
                try:
                    dt = xldate.xldate_as_datetime(v, datemode)
                    if dt.hour == 0 and dt.minute == 0 and dt.second == 0:
                        return dt.strftime("%Y-%m-%d")
                    return dt.strftime("%Y-%m-%d %H:%M:%S")
                except Exception:
                    return v
            return v

        # 修剪尾部纯空的行/列
        while nrows > 0 and all(_xls_val(nrows - 1, c) in (None, "") for c in range(ncols)):
            nrows -= 1
        while ncols > 0 and all(_xls_val(r, ncols - 1) in (None, "") for r in range(nrows)):
            ncols -= 1

        if nrows == 0 or ncols == 0:
            print("_empty sheet_")
            print()
            continue

        header = "| " + " | ".join(["_row/col_"] + [_cell_col_letter(c + 1) for c in range(ncols)]) + " |"
        sep = "| " + " | ".join(["---"] * (ncols + 1)) + " |"
        print(header)
        print(sep)

        for r in range(nrows):
            row_cells = [str(r + 1)]
            for c in range(ncols):
                v = _xls_val(r, c)
                # xlrd 的 float 需要处理成整数字符串（若是整数值）
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                row_cells.append(_clean_cell(v).replace("|", "\\|"))
            print("| " + " | ".join(row_cells) + " |")
        print()
        if merged_ranges:
            print("### Merged ranges")
            print()
            for merged_range in merged_ranges:
                print(f"- `{merged_range}`")
            print()

    return {
        "coverage": {
            "sheet_count": book.nsheets,
            "merged_ranges": sum(len(sheet.merged_cells) for sheet in book.sheets()),
            "formula_inspection_supported": False,
            "complete": False,
            "omissions": ["legacy_xls_formula_inspection_unavailable"],
        },
        "issues": [
            {
                "code": "legacy_xls_formula_unverified",
                "message": "旧版 XLS 无法可靠读取公式定义，必须对照原文件复核",
                "source_values": [],
            }
        ],
    }


# ---------- 分支：docx ----------


def _word_paragraph_text(
    paragraph: ElementTree.Element,
    parent_map: dict[ElementTree.Element, ElementTree.Element],
) -> str:
    """Read text owned by one paragraph without consuming nested text boxes."""
    parts: list[str] = []
    for node in paragraph.iter():
        name = _xml_local_name(node.tag)
        if name not in {"t", "tab", "br", "cr"}:
            continue
        owner = parent_map.get(node)
        while owner is not None and _xml_local_name(owner.tag) != "p":
            owner = parent_map.get(owner)
        if owner is not paragraph:
            continue
        if name == "t" and node.text:
            parts.append(node.text)
        elif name == "tab":
            parts.append("\t")
        else:
            parts.append("\n")
    return "".join(parts).strip()


def _has_word_ancestor(
    element: ElementTree.Element,
    parent_map: dict[ElementTree.Element, ElementTree.Element],
    local_name: str,
) -> bool:
    parent = parent_map.get(element)
    while parent is not None:
        if _xml_local_name(parent.tag) == local_name:
            return True
        parent = parent_map.get(parent)
    return False


def _word_relationships_path(part_name: str) -> str:
    directory, filename = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _inspect_docx_package(data: bytes) -> dict[str, object]:
    """Extract OOXML story text, text boxes, and embedded image assets."""
    story_names: list[str] = []
    story_sections: list[tuple[str, list[str]]] = []
    text_boxes: list[tuple[str, str]] = []
    embedded_images: list[dict[str, object]] = []
    omissions: list[str] = []
    converted_parts = 0

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
        story_names = sorted(
            name
            for name in names
            if name == "word/document.xml"
            or name.startswith("word/header") and name.endswith(".xml")
            or name.startswith("word/footer") and name.endswith(".xml")
            or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
        )
        seen_image_refs: set[tuple[str, str]] = set()
        for part_name in story_names:
            try:
                root = ElementTree.fromstring(archive.read(part_name))
            except (ElementTree.ParseError, KeyError):
                omissions.append(part_name)
                continue
            converted_parts += 1
            parent_map = {child: parent for parent in root.iter() for child in parent}

            normal_paragraphs: list[str] = []
            for paragraph in (node for node in root.iter() if _xml_local_name(node.tag) == "p"):
                in_text_box = _has_word_ancestor(paragraph, parent_map, "txbxContent")
                value = _word_paragraph_text(paragraph, parent_map)
                if in_text_box:
                    if value:
                        text_boxes.append((part_name, value))
                elif part_name != "word/document.xml" and value:
                    normal_paragraphs.append(value)
            if normal_paragraphs:
                story_sections.append((part_name, normal_paragraphs))

            relationships: dict[str, str] = {}
            rels_name = _word_relationships_path(part_name)
            if rels_name in names:
                try:
                    rels_root = ElementTree.fromstring(archive.read(rels_name))
                    relationships = {
                        str(rel.attrib.get("Id")): str(rel.attrib.get("Target"))
                        for rel in rels_root
                        if rel.attrib.get("Id") and rel.attrib.get("Target")
                    }
                except ElementTree.ParseError:
                    omissions.append(rels_name)

            for node in root.iter():
                if _xml_local_name(node.tag) not in {"blip", "imagedata"}:
                    continue
                relationship_id = next(
                    (
                        value
                        for key, value in node.attrib.items()
                        if _xml_local_name(key) in {"embed", "id"}
                    ),
                    None,
                )
                if not relationship_id or (part_name, relationship_id) in seen_image_refs:
                    continue
                seen_image_refs.add((part_name, relationship_id))
                target = relationships.get(relationship_id)
                if not target or "://" in target:
                    omissions.append(f"{part_name}#{relationship_id}")
                    continue
                media_name = posixpath.normpath(posixpath.join(posixpath.dirname(part_name), target))
                if media_name not in names:
                    omissions.append(media_name)
                    continue
                embedded_images.append(
                    {
                        "part": part_name,
                        "relationship_id": relationship_id,
                        "filename": posixpath.basename(media_name),
                        "bytes": archive.read(media_name),
                    }
                )

    return {
        "story_sections": story_sections,
        "text_boxes": text_boxes,
        "embedded_images": embedded_images,
        "coverage": {
            "story_parts": len(story_names),
            "story_parts_converted": converted_parts,
            "text_boxes": len(text_boxes),
            "embedded_images": len(embedded_images),
            "embedded_images_processed": 0,
            "complete": not omissions,
            "omissions": omissions,
        },
    }


def convert_docx(path: str) -> dict[str, object]:
    from docx import Document
    from docx.oxml.ns import qn

    data = Path(path).read_bytes()
    package_report = _inspect_docx_package(data)
    doc = Document(path)
    print(f"# {Path(path).name}")
    print("_format: docx_")
    print()

    # 按 body 顺序遍历段落和表格
    body = doc.element.body
    p_iter = iter(doc.paragraphs)
    t_iter = iter(doc.tables)

    p_idx = 0
    t_idx = 0
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            para = next(p_iter, None)
            if para is None:
                continue
            text = para.text or ""
            text = text.replace("\r\n", "\n").replace("\r", "\n").rstrip()
            p_idx += 1
            if text.strip() == "":
                # 保留空行为一个明确的占位，便于人工/LLM 定位段落断点
                print(f"_p{p_idx}: (empty)_")
            elif "\n" in text:
                # 段内软换行（shift-enter）：每个子行独立成 markdown 行，便于 LLM 逐字段抽取
                lines = [ln for ln in text.split("\n")]
                for sub_i, sub in enumerate(lines, start=1):
                    if sub.strip() == "":
                        print(f"_p{p_idx}.{sub_i}: (empty)_")
                    else:
                        print(f"_p{p_idx}.{sub_i}_ {sub.rstrip()}")
            else:
                print(f"_p{p_idx}_ {text}")
            print()
        elif tag == qn("w:tbl"):
            tbl = next(t_iter, None)
            if tbl is None:
                continue
            t_idx += 1
            print(f"### Table {t_idx}")
            print()
            rows = tbl.rows
            if not rows:
                print("_empty table_")
                print()
                continue
            ncols = max(len(r.cells) for r in rows)
            header = "| " + " | ".join(["_row/col_"] + [_cell_col_letter(c + 1) for c in range(ncols)]) + " |"
            sep = "| " + " | ".join(["---"] * (ncols + 1)) + " |"
            print(header)
            print(sep)
            seen_table_cells: set[int] = set()
            for r_idx, row in enumerate(rows, start=1):
                row_cells = [str(r_idx)]
                cells = row.cells
                for c_idx in range(ncols):
                    if c_idx < len(cells):
                        cell_key = id(cells[c_idx]._tc)
                        if cell_key in seen_table_cells:
                            v = ""
                        else:
                            seen_table_cells.add(cell_key)
                            v = cells[c_idx].text or ""
                    else:
                        v = ""
                    row_cells.append(_clean_cell(v).replace("|", "\\|"))
                print("| " + " | ".join(row_cells) + " |")
            print()

    for part_name, paragraphs in package_report["story_sections"]:
        print(f"## Story: {part_name}")
        print()
        for paragraph_index, text in enumerate(paragraphs, 1):
            print(f"_p{paragraph_index}_ {text}")
            print()

    text_boxes = package_report["text_boxes"]
    if text_boxes:
        print("## Text boxes")
        print()
        for text_box_index, (part_name, text) in enumerate(text_boxes, 1):
            print(f"### Text box {text_box_index}")
            print(f"_source: {part_name}_")
            print()
            print(text)
            print()

    embedded_images = package_report["embedded_images"]
    if embedded_images:
        print("## Embedded images")
        print()
        for image_index, image in enumerate(embedded_images, 1):
            print(f"### Image {image_index}: {image['filename']}")
            print(
                f"_source: {image['part']}#{image['relationship_id']}; "
                "content is routed separately for image verification_"
            )
            print()

    return package_report


# ---------- 分支：doc（需 libreoffice） ----------

def _find_soffice() -> str | None:
    """跨平台定位 LibreOffice 可执行文件。"""
    # PATH 命中优先（Linux/Docker 常见）
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    # 常见平台安装位置兜底
    candidates = [
        # macOS
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        # Linux 常见路径
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
        "/usr/lib/libreoffice/program/soffice",
        "/opt/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
        # Windows
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return None


def _find_textutil() -> str | None:
    """Locate macOS' built-in Word converter as a local fallback."""
    return shutil.which("textutil")


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _is_word_xml(path: str) -> bool:
    try:
        prefix = Path(path).read_bytes()[:4096].decode("utf-8-sig", errors="ignore")
    except OSError:
        return False
    compact = prefix.lstrip()
    return compact.startswith("<?xml") and (
        "schemas.microsoft.com/office/word/2003/wordml" in prefix
        or "schemas.microsoft.com/office/2006/xmlPackage" in prefix
    )


def _word_xml_text(element: ElementTree.Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        name = _xml_local_name(node.tag)
        if name == "t" and node.text:
            parts.append(node.text)
        elif name == "tab":
            parts.append("\t")
        elif name in {"br", "cr"}:
            parts.append("\n")
    return "".join(parts).strip()


def convert_word_xml(path: str) -> None:
    """Convert Word 2003 XML or Flat OPC XML without external office tools."""
    root = ElementTree.parse(path).getroot()
    body = next((node for node in root.iter() if _xml_local_name(node.tag) == "body"), None)
    if body is None:
        raise ValueError("Word XML document has no body")

    print(f"# {Path(path).name}")
    print("_format: word_xml_")
    print()
    paragraph_index = 0
    table_index = 0
    for child in body:
        name = _xml_local_name(child.tag)
        if name == "p":
            paragraph_index += 1
            value = _word_xml_text(child)
            print(f"_p{paragraph_index}_ {value}" if value else f"_p{paragraph_index}: (empty)_")
            print()
            continue
        if name != "tbl":
            continue

        table_index += 1
        rows = [node for node in child if _xml_local_name(node.tag) == "tr"]
        print(f"### Table {table_index}")
        print()
        if not rows:
            print("_empty table_")
            print()
            continue
        parsed_rows = [
            [_word_xml_text(cell) for cell in row if _xml_local_name(cell.tag) == "tc"]
            for row in rows
        ]
        column_count = max((len(row) for row in parsed_rows), default=0)
        header = "| " + " | ".join(
            ["_row/col_", *[_cell_col_letter(index + 1) for index in range(column_count)]]
        ) + " |"
        print(header)
        print("| " + " | ".join(["---"] * (column_count + 1)) + " |")
        for row_index, row in enumerate(parsed_rows, start=1):
            cells = [str(row_index)]
            cells.extend(_clean_cell(value).replace("|", "\\|") for value in row)
            cells.extend([""] * (column_count - len(row)))
            print("| " + " | ".join(cells) + " |")
        print()


def convert_doc(path: str) -> dict[str, object] | None:
    if _is_word_xml(path):
        convert_word_xml(path)
        return

    soffice = _find_soffice()
    textutil = _find_textutil() if not soffice else None
    if not soffice and not textutil:
        emit_scan_hint(
            path,
            reason=(
                "doc 需 LibreOffice 或 macOS textutil 转换（均未找到）；"
                "请安装 LibreOffice 或改用 OCR"
            ),
        )
        return

    with tempfile.TemporaryDirectory() as tmpd:
        stem = Path(path).stem
        docx_path = Path(tmpd) / f"{stem}.docx"
        if soffice:
            command = [soffice, "--headless", "--convert-to", "docx", "--outdir", tmpd, path]
            converter_name = "libreoffice"
        else:
            command = [textutil, "-convert", "docx", "-output", str(docx_path), "--", path]
            converter_name = "textutil"
        try:
            subprocess.run(
                command,
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
                timeout=120,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            emit_scan_hint(path, reason=f"{converter_name} 转换失败: {e}")
            return

        if not docx_path.exists():
            # libreoffice 有时会改名，回退用目录里唯一的 docx
            candidates = list(Path(tmpd).glob("*.docx"))
            if not candidates:
                emit_scan_hint(path, reason="libreoffice 未产出 docx")
                return
            docx_path = candidates[0]

        return convert_docx(str(docx_path))


# ---------- 分支：pdf ----------

def convert_pdf(path: str) -> None:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        total_text_len = 0
        for page in pdf.pages:
            t = page.extract_text() or ""
            total_text_len += len(t.strip())

        # 全文极短 → 扫描件，交 OCR
        if total_text_len < 40:
            emit_scan_hint(path, reason=f"pdf 文本层过短 ({total_text_len} 字符)，判定为扫描件")
            return

        print(f"# {Path(path).name}")
        print(f"_format: pdf_ pages: {len(pdf.pages)}")
        print()

        for page_idx, page in enumerate(pdf.pages, start=1):
            print(f"## Page {page_idx}")
            print()

            # 文本层
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                print("### Text")
                print()
                for line_i, line in enumerate(text.split("\n"), start=1):
                    print(f"_l{line_i}_ {line.rstrip()}")
                print()
            else:
                print("_no text layer on this page_")
                print()

            # 表格层（若有）
            tables = []
            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []

            for t_idx, table in enumerate(tables, start=1):
                print(f"### Table {page_idx}.{t_idx}")
                print()
                if not table:
                    print("_empty table_")
                    print()
                    continue
                ncols = max(len(r) for r in table)
                header = "| " + " | ".join(["_row/col_"] + [_cell_col_letter(c + 1) for c in range(ncols)]) + " |"
                sep = "| " + " | ".join(["---"] * (ncols + 1)) + " |"
                print(header)
                print(sep)
                for r_idx, row in enumerate(table, start=1):
                    cells = [str(r_idx)]
                    for c_idx in range(ncols):
                        v = row[c_idx] if c_idx < len(row) else ""
                        cells.append(_clean_cell(v).replace("|", "\\|"))
                    print("| " + " | ".join(cells) + " |")
                print()


# ---------- 分支：图片 ----------

def convert_image(path: str) -> None:
    emit_scan_hint(path, reason="image → OCR")


# ---------- 入口 ----------

DISPATCH = {
    ".xlsx": convert_xlsx,
    ".xlsm": convert_xlsx,
    ".xls": convert_xls,
    ".docx": convert_docx,
    ".doc": convert_doc,
    ".pdf": convert_pdf,
}


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("usage: to_text.py <file>", file=sys.stderr)
        return 2

    path = argv[1]
    if not os.path.exists(path):
        print(f"file not found: {path}", file=sys.stderr)
        return 2

    ext = Path(path).suffix.lower()

    if ext in IMAGE_EXTS:
        convert_image(path)
        return 0

    handler = DISPATCH.get(ext)
    if handler is None:
        print(f"unsupported extension: {ext}", file=sys.stderr)
        return 2

    try:
        handler(path)
    except Exception as e:
        print(f"convert error: {e.__class__.__name__}: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
